"""Validation and parsing helpers for Key Vault bulk secret operations."""

from __future__ import annotations

import csv
import io
import json
import re
from typing import Any

SECRET_NAME_PATTERN = re.compile(r"^[a-zA-Z0-9-]+$")
SECRET_NAME_MAX_LEN = 127
SECRET_VALUE_MAX_BYTES = 25 * 1024  # Azure Key Vault limit
BULK_SECRET_MAX_COUNT = 1000


def validate_bulk_secrets(
    secrets: list[dict[str, Any]],
    *,
    vault_uri: str,
) -> dict[str, Any]:
    """Validate a batch of secrets before upload. Never logs secret values."""
    if not vault_uri or not vault_uri.strip():
        return {
            "valid": False,
            "total": len(secrets),
            "errors": [{"row": 0, "name": "", "error": "Target Key Vault is required"}],
            "items": [],
        }

    if len(secrets) > BULK_SECRET_MAX_COUNT:
        return {
            "valid": False,
            "total": len(secrets),
            "errors": [
                {
                    "row": 0,
                    "name": "",
                    "error": f"Maximum {BULK_SECRET_MAX_COUNT} secrets per bulk upload",
                }
            ],
            "items": [],
        }

    seen_names: dict[str, int] = {}
    errors: list[dict[str, Any]] = []
    items: list[dict[str, Any]] = []

    for index, secret in enumerate(secrets, start=1):
        name = (secret.get("name") or "").strip()
        value = secret.get("value")
        row_errors: list[str] = []

        if not name:
            row_errors.append("Secret name is required")
        elif len(name) > SECRET_NAME_MAX_LEN:
            row_errors.append(f"Secret name exceeds {SECRET_NAME_MAX_LEN} characters")
        elif not SECRET_NAME_PATTERN.match(name):
            row_errors.append("Secret name must be alphanumeric and hyphens only")
        elif name in seen_names:
            row_errors.append(f"Duplicate secret name (also on row {seen_names[name]})")
        else:
            seen_names[name] = index

        if value is None or (isinstance(value, str) and not value.strip()):
            row_errors.append("Secret value is required")
        elif isinstance(value, str):
            value_bytes = value.encode("utf-8")
            if len(value_bytes) > SECRET_VALUE_MAX_BYTES:
                row_errors.append(f"Secret value exceeds {SECRET_VALUE_MAX_BYTES} bytes")

        if row_errors:
            for message in row_errors:
                errors.append({"row": index, "name": name or f"row-{index}", "error": message})
            items.append({"row": index, "name": name, "status": "invalid", "errors": row_errors})
        else:
            items.append({"row": index, "name": name, "status": "valid", "errors": []})

    return {
        "valid": len(errors) == 0,
        "total": len(secrets),
        "valid_count": sum(1 for item in items if item["status"] == "valid"),
        "invalid_count": sum(1 for item in items if item["status"] == "invalid"),
        "errors": errors,
        "items": items,
    }


def _normalize_row(row: dict[str, Any]) -> dict[str, Any]:
    """Map flexible column names to canonical secret fields."""
    key_map = {
        "name": "name",
        "secret_name": "name",
        "secretname": "name",
        "value": "value",
        "secret_value": "value",
        "secretvalue": "value",
        "content_type": "content_type",
        "contenttype": "content_type",
        "expires": "expires",
        "expiry": "expires",
        "expiry_date": "expires",
        "tags": "tags",
    }
    normalized: dict[str, Any] = {}
    for raw_key, raw_value in row.items():
        if raw_value is None:
            continue
        key = key_map.get(str(raw_key).strip().lower().replace(" ", "_"))
        if not key:
            continue
        if key == "tags" and isinstance(raw_value, str) and raw_value.strip():
            try:
                normalized[key] = json.loads(raw_value)
            except json.JSONDecodeError:
                normalized[key] = {
                    pair.split("=", 1)[0].strip(): pair.split("=", 1)[1].strip()
                    for pair in raw_value.split(",")
                    if "=" in pair
                }
        else:
            normalized[key] = str(raw_value).strip() if key != "tags" else raw_value
    return normalized


def parse_bulk_secrets_csv(content: bytes) -> list[dict[str, Any]]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return [_normalize_row(row) for row in reader]


def parse_bulk_secrets_json(content: bytes) -> list[dict[str, Any]]:
    data = json.loads(content.decode("utf-8"))
    if isinstance(data, dict) and "secrets" in data:
        rows = data["secrets"]
    elif isinstance(data, list):
        rows = data
    else:
        raise ValueError("JSON must be an array or object with a 'secrets' array")
    return [_normalize_row(row) if isinstance(row, dict) else {} for row in rows]


def parse_bulk_secrets_xlsx(content: bytes) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    if sheet is None:
        raise ValueError("Excel file has no active worksheet")

    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return []

    headers = [str(cell).strip() if cell is not None else "" for cell in header_row]
    secrets: list[dict[str, Any]] = []
    for row in rows_iter:
        if not any(cell is not None and str(cell).strip() for cell in row):
            continue
        row_dict = {headers[i]: row[i] for i in range(min(len(headers), len(row))) if headers[i]}
        secrets.append(_normalize_row(row_dict))
    return secrets


def parse_bulk_secrets_file(filename: str, content: bytes) -> list[dict[str, Any]]:
    """Parse CSV, JSON, or XLSX bulk secret upload files."""
    lower = filename.lower()
    if lower.endswith(".csv"):
        return parse_bulk_secrets_csv(content)
    if lower.endswith(".json"):
        return parse_bulk_secrets_json(content)
    if lower.endswith(".xlsx"):
        return parse_bulk_secrets_xlsx(content)
    raise ValueError("Unsupported file type. Use .csv, .json, or .xlsx")
