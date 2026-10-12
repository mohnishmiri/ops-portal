"""Credential account inventory — the service accounts the team owns, by application.

The Credential Check answers "does this credential work right now?"; this
keeps the list it is asked about: who owns each domain account, AAF identity
and JFrog user, when it expires, when it last signed in, how many failed
logins it has and where to reset it.

**No password or token is ever stored here.**  The values come from three
places, and only these:

* the Add / Edit form, which replaces every field;
* an Excel / CSV import, which fills in and corrects but never erases — a
  blank cell keeps what is already recorded;
* a successful Credential Check of an account that is listed, which records
  the verdict and whatever the source reported (expiry, password change,
  last sign-in).  Failed logins are never taken from a check: a check's own
  failures are not the account's sign-in history.

Every change made from the page is written to ``audit_logs``.  A check's
update of a row is not — the row itself records who checked it and when, and
checks are deliberately not logged.
"""

from __future__ import annotations

import csv
import io
import re
from collections.abc import Callable, Iterable
from dataclasses import dataclass
from datetime import UTC, date, datetime, timedelta
from typing import Any

import structlog
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.auth import UserContext
from app.models.database import AuditLog, CredentialAccount
from app.schemas.credential_inventory import (
    DIRECTORIES,
    CredentialAccountImportRow,
    CredentialAccountIn,
    ImportApplyResult,
    ImportPreview,
    ImportPreviewRow,
)
from app.services.compliance_excel_service import (
    _apply_freeze_and_filter,
    _auto_fit_columns,
    _write_branding_header,
    _write_column_headers,
    _write_data_cell,
)
from app.services.credential_check_service import CheckOutcome

logger = structlog.get_logger(__name__)

IMPORT_MAX_ROWS = 2000
IMPORT_MAX_BYTES = 5 * 1024 * 1024

TYPE_LABELS = {"domain": "Domain", "aaf": "AAF", "jfrog": "JFrog"}
DIRECTORY_LABELS = {
    "ITSERVICES": "ITSERVICES",
    "TESTITSERVICES": "TESTITSERVICES",
    "prod": "Prod",
    "nonprod": "IST",
    "bronze": "Bronze",
    "gold": "Gold",
}

# A last check with one of these verdicts means the account cannot sign in.
FAILING_STATUSES = frozenset({"invalid", "expired", "locked", "disabled", "must_change", "denied"})
EXPIRING_WINDOW_DAYS = 30

_AUDIT_RESOURCE_TYPE = "credential_account"


class DuplicateAccountError(Exception):
    """The account is already listed."""


class _NoHeaderRowError(ValueError):
    """A sheet with no recognisable header — the workbook may have another."""


def account_label(account: Any) -> str:
    """ "ITSERVICES\\m52142", "AAF Prod: m1@com.att.attcc", "JFrog Gold: m52142@att.com"."""
    if account.account_type == "domain":
        return f"{account.directory}\\{account.account_name}"
    prefix = f"{TYPE_LABELS[account.account_type]} {DIRECTORY_LABELS.get(account.directory, account.directory)}"
    return f"{prefix}: {account.account_name}"


def _utc_date(value: datetime) -> date:
    return (value.astimezone(UTC) if value.tzinfo else value).date()


# ── Import: reading a spreadsheet ──────────────────────────────────────

# Header text (lower-cased, spaces / underscores collapsed) → field.  Generous
# on purpose: people's mechID sheets are not written to a template.
_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "application": ("application", "application name", "app", "app name"),
    "app_code": ("app code", "mots id", "motsid", "mots"),
    "account_type": ("type", "account type"),
    "directory": ("directory", "domain", "instance", "environment", "env"),
    "account_name": (
        "account",
        "account name",
        "mechid",
        "mech id",
        "username",
        "user name",
        "user",
        "identity",
        "login",
    ),
    "owner": ("owner", "owner name", "owner email", "account owner", "owned by"),
    "expiry_date": (
        "expiry",
        "expiry date",
        "expires",
        "expiration",
        "expiration date",
        "password expiry",
        "password expires",
    ),
    "password_changed_date": (
        "password last changed",
        "password changed",
        "last changed",
        "pwd last set",
        "password last set",
    ),
    "last_login_date": ("last login", "last login date", "last logon", "last sign-in", "last sign in", "last used"),
    "failed_logins": ("failed logins", "failed login", "failed login count", "failed sign-ins", "bad password count"),
    "reset_url": ("reset link", "reset url", "reset", "password reset link", "reset password link"),
    "notes": ("notes", "additional information", "additional info", "comments", "comment", "description"),
}
_HEADER_LOOKUP = {alias: field for field, aliases in _HEADER_ALIASES.items() for alias in aliases}

FIELD_LABELS = {
    "application": "Application",
    "app_code": "App Code",
    "account_type": "Type",
    "directory": "Directory",
    "account_name": "Account",
    "owner": "Owner",
    "expiry_date": "Expiry Date",
    "password_changed_date": "Password Last Changed",
    "last_login_date": "Last Login",
    "failed_logins": "Failed Logins",
    "reset_url": "Reset Link",
    "notes": "Notes",
}

# "Type" cell words.  ITSERVICES / TESTITSERVICES name the directory too.
_TYPE_WORDS: dict[str, tuple[str, str | None]] = {
    "itservices": ("domain", "ITSERVICES"),
    "testitservices": ("domain", "TESTITSERVICES"),
    "domain": ("domain", None),
    "ad": ("domain", None),
    "active directory": ("domain", None),
    "aaf": ("aaf", None),
    "jfrog": ("jfrog", None),
    "artifactory": ("jfrog", None),
}
_DIRECTORY_WORDS: dict[str, str] = {
    "itservices": "ITSERVICES",
    "testitservices": "TESTITSERVICES",
    "prod": "prod",
    "production": "prod",
    "ist": "nonprod",
    "nonprod": "nonprod",
    "non-prod": "nonprod",
    "non prod": "nonprod",
    "bronze": "bronze",
    "artifact.it.att.com": "bronze",
    "gold": "gold",
    "artifact-gold.it.att.com": "gold",
}

_NEVER_WORDS = frozenset({"never", "does not expire", "doesn't expire", "no expiry", "non-expiring"})
_BLANK_WORDS = frozenset({"", "-", "—", "n/a", "na", "none", "null"})
_DATE_FORMATS = ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%m-%d-%Y", "%d-%b-%Y", "%d %b %Y", "%b %d, %Y", "%d-%b-%y")


def _norm(value: Any) -> str:
    return re.sub(r"[\s_]+", " ", str(value or "")).strip().lower()


def _cell_text(value: Any) -> str | None:
    """A cell as text; Excel's 31599.0 becomes "31599"."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, datetime | date):
        return value.isoformat()
    text = str(value).strip()
    return text or None


def _parse_date(value: Any) -> tuple[date | None, bool]:
    """(date, says "never").  Raises ValueError for anything unreadable."""
    if value is None:
        return None, False
    if isinstance(value, datetime):
        return value.date(), False
    if isinstance(value, date):
        return value, False
    if isinstance(value, int | float) and not isinstance(value, bool):
        # An Excel date serial in a cell that was not formatted as a date.
        if 20_000 < value < 80_000:
            return date(1899, 12, 30) + timedelta(days=int(value)), False
        raise ValueError(f"'{value}' is not a date")
    text = str(value).strip()
    lowered = text.lower()
    if lowered in _BLANK_WORDS:
        return None, False
    if lowered in _NEVER_WORDS:
        return None, True
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date(), False
    except ValueError:
        pass
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date(), False
        except ValueError:
            continue
    raise ValueError(f"'{text}' is not a date — use YYYY-MM-DD or MM/DD/YYYY")


def _parse_count(value: Any) -> int | None:
    if value is None or (isinstance(value, str) and value.strip().lower() in _BLANK_WORDS):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        raise ValueError(f"'{value}' is not a number") from None
    if not number.is_integer() or number < 0:
        raise ValueError(f"'{value}' is not a whole number")
    return int(number)


def _resolve_type(type_text: str | None, directory_text: str | None) -> tuple[str, str]:
    """("domain", "ITSERVICES") from cells such as "ITSERVICES", "AAF" + "IST", "JFrog Gold"."""
    words = _norm(type_text)
    account_type, directory = _TYPE_WORDS.get(words, (None, None))
    if account_type is None:
        # "JFrog Gold", "AAF - Non-Prod": a type word plus a directory phrase.
        # Phrases longest first, whole words only, so "non-prod" is not "prod".
        account_type, directory = next(
            (_TYPE_WORDS[word] for word in re.split(r"[\s\-/:(),]+", words) if word in _TYPE_WORDS),
            (None, None),
        )
        if account_type is not None and directory is None:
            directory = next(
                (
                    value
                    for phrase, value in sorted(_DIRECTORY_WORDS.items(), key=lambda item: -len(item[0]))
                    if re.search(rf"(?<![\w.-]){re.escape(phrase)}(?![\w.-])", words)
                ),
                None,
            )
    named = None
    if directory_text and _norm(directory_text):
        named = _DIRECTORY_WORDS.get(_norm(directory_text))
        if named is None:
            raise ValueError(f"Directory {directory_text.strip()!r} is not recognised")
    if account_type is None and not words and named:
        # No Type column, but "ITSERVICES" / "Gold" / "IST" can only be one kind.
        account_type = next(kind for kind, allowed in DIRECTORIES.items() if named in allowed)
    if account_type is None:
        shown = (type_text or "").strip() or "(blank)"
        raise ValueError(f"Type {shown!r} is not one of ITSERVICES, TESTITSERVICES, AAF or JFrog")

    if named:
        directory = named
    directory = directory or DIRECTORIES[account_type][0]
    if directory not in DIRECTORIES[account_type]:
        raise ValueError(
            f"Directory {DIRECTORY_LABELS.get(directory, directory)!r} does not fit a "
            f"{TYPE_LABELS[account_type]} account"
        )
    return account_type, directory


def _header_map(row: Iterable[Any]) -> dict[int, str]:
    """Column index → field for a header row; {} when the row is not a header."""
    mapping: dict[int, str] = {}
    for index, cell in enumerate(row):
        field = _HEADER_LOOKUP.get(_norm(cell))
        if field and field not in mapping.values():
            mapping[index] = field
    return mapping if "account_name" in mapping.values() else {}


def _table_rows(raw_rows: Iterable[tuple[int, list[Any]]]) -> list[tuple[int, dict[str, Any]]]:
    """Find the header (within the first rows — an export has a title above it) and map the rest."""
    header: dict[int, str] = {}
    rows: list[tuple[int, dict[str, Any]]] = []
    for number, cells in raw_rows:
        if not header:
            if number > 15:
                break
            header = _header_map(cells)
            continue
        if not any(cell is not None and str(cell).strip() for cell in cells):
            continue
        if len(rows) >= IMPORT_MAX_ROWS:
            raise ValueError(f"The file lists more than {IMPORT_MAX_ROWS} accounts — split it into smaller files.")
        rows.append((number, {field: cells[index] for index, field in header.items() if index < len(cells)}))
    if not header:
        raise _NoHeaderRowError(
            "No header row found. The first row should name the columns — at least Application, Type and Account. "
            "Download the template to see the expected layout."
        )
    return rows


def _read_xlsx(content: bytes) -> list[tuple[int, dict[str, Any]]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        last_error: ValueError = ValueError("The workbook has no worksheets.")
        # The template's instructions sheet has no Account column, so the
        # first sheet with a recognisable header is the data.
        for sheet in workbook.worksheets:
            try:
                return _table_rows(
                    (number, list(cells)) for number, cells in enumerate(sheet.iter_rows(values_only=True), start=1)
                )
            except _NoHeaderRowError as exc:
                last_error = exc
        raise last_error
    finally:
        workbook.close()


def _read_csv(content: bytes) -> list[tuple[int, dict[str, Any]]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Excel's "CSV" on Windows saves in the ANSI code page, not UTF-8.
        text = content.decode("cp1252", errors="replace")
    reader = csv.reader(io.StringIO(text))
    return _table_rows((number, list(cells)) for number, cells in enumerate(reader, start=1))


def read_import_file(filename: str, content: bytes) -> list[tuple[int, dict[str, Any]]]:
    """(spreadsheet row number, {field: raw cell}) for every non-blank data row."""
    lower = (filename or "").lower()
    if lower.endswith(".xlsx"):
        try:
            return _read_xlsx(content)
        except ValueError:
            raise
        except Exception as exc:  # a corrupt or non-Excel file
            raise ValueError("The file could not be read as an Excel workbook (.xlsx).") from exc
    if lower.endswith(".csv"):
        return _read_csv(content)
    if lower.endswith(".xls"):
        raise ValueError("Old-format .xls files are not supported — save it as .xlsx or .csv and upload again.")
    raise ValueError("Upload an Excel (.xlsx) or CSV file.")


def _pydantic_errors(exc: ValidationError) -> list[str]:
    messages = []
    for error in exc.errors():
        field = error["loc"][0] if error.get("loc") else None
        message = str(error.get("msg", "")).removeprefix("Value error, ")
        messages.append(f"{FIELD_LABELS.get(str(field), 'Row')}: {message}" if field else message)
    return messages


def parse_import_row(raw: dict[str, Any]) -> tuple[CredentialAccountImportRow | None, list[str]]:
    """One spreadsheet row → an import row, or the reasons it cannot be one."""
    errors: list[str] = []
    values: dict[str, Any] = {}

    try:
        values["account_type"], values["directory"] = _resolve_type(
            _cell_text(raw.get("account_type")), _cell_text(raw.get("directory"))
        )
    except ValueError as exc:
        errors.append(str(exc))

    for field in ("application", "app_code", "account_name", "owner", "reset_url", "notes"):
        values[field] = _cell_text(raw.get(field))

    for field in ("expiry_date", "password_changed_date", "last_login_date"):
        try:
            parsed, never = _parse_date(raw.get(field))
        except ValueError as exc:
            errors.append(f"{FIELD_LABELS[field]}: {exc}")
            continue
        values[field] = parsed
        if field == "expiry_date" and (parsed or never):
            values["never_expires"] = never

    try:
        values["failed_logins"] = _parse_count(raw.get("failed_logins"))
    except ValueError as exc:
        errors.append(f"{FIELD_LABELS['failed_logins']}: {exc}")

    if not values.get("application"):
        errors.append("Application: required")
    if not values.get("account_name"):
        errors.append("Account: required")
    if errors:
        return None, errors

    try:
        return CredentialAccountImportRow(**values), []
    except ValidationError as exc:
        return None, _pydantic_errors(exc)


# ── Export ─────────────────────────────────────────────────────────────


def _days_left(account: CredentialAccount, today: date) -> int | None:
    return (account.expiry_date - today).days if account.expiry_date and not account.never_expires else None


def _export_columns(today: date) -> list[tuple[str, Callable[[CredentialAccount], Any]]]:
    """Same headers the import reads, so an export can be edited and imported back."""
    return [
        ("Application", lambda a: a.application),
        ("App Code", lambda a: a.app_code),
        ("Type", lambda a: TYPE_LABELS[a.account_type]),
        ("Directory", lambda a: DIRECTORY_LABELS.get(a.directory, a.directory)),
        ("Account", lambda a: a.account_name),
        ("Directory Name", lambda a: a.display_name),
        ("Owner", lambda a: a.owner),
        ("Expiry Date", lambda a: "Never" if a.never_expires else a.expiry_date),
        ("Days Left", lambda a: _days_left(a, today)),
        ("Password Last Changed", lambda a: a.password_changed_date),
        ("Last Login", lambda a: a.last_login_date),
        ("Failed Logins", lambda a: a.failed_logins),
        ("Last Check", lambda a: a.last_check_status),
        ("Last Checked (UTC)", lambda a: a.last_checked_at.replace(microsecond=0) if a.last_checked_at else None),
        ("Reset Link", lambda a: a.reset_url),
        ("Notes", lambda a: a.notes),
    ]


def _csv_safe(value: Any) -> Any:
    """Neutralise spreadsheet formulas: a cell starting =, +, - or @ is text, not a formula."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def export_csv(accounts: list[CredentialAccount], *, today: date) -> bytes:
    columns = _export_columns(today)
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow([header for header, _ in columns])
    for account in accounts:
        row = []
        for _header, value in columns:
            cell = value(account)
            row.append("" if cell is None else _csv_safe(cell.isoformat() if isinstance(cell, date) else cell))
        writer.writerow(row)
    # BOM so Excel opens it as UTF-8.
    return ("﻿" + buffer.getvalue()).encode("utf-8")


def _write_text_safe(ws: Any, row: int, col: int, value: Any) -> None:
    _write_data_cell(ws, row, col, value)
    cell = ws.cell(row=row, column=col)
    if isinstance(value, str) and value.startswith("="):
        # openpyxl stores "=..." as a formula; an imported note must stay text.
        cell.data_type = "s"
    elif isinstance(value, datetime):
        cell.number_format = "yyyy-mm-dd hh:mm"
    elif isinstance(value, date):
        cell.number_format = "yyyy-mm-dd"


def export_xlsx(accounts: list[CredentialAccount], *, today: date, generated_at: datetime) -> bytes:
    columns = _export_columns(today)
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Accounts"
    header_row = _write_branding_header(ws, "Credential account inventory", generated_at)
    _write_column_headers(ws, header_row, [header for header, _ in columns])
    for offset, account in enumerate(accounts, start=1):
        for col, (_header, value) in enumerate(columns, start=1):
            _write_text_safe(ws, header_row + offset, col, value(account))
    last_row = header_row + max(len(accounts), 1)
    _auto_fit_columns(ws, header_row, last_row)
    _apply_freeze_and_filter(ws, header_row, len(columns))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


TEMPLATE_HEADERS = [
    "Application",
    "App Code",
    "Type",
    "Directory",
    "Account",
    "Owner",
    "Expiry Date",
    "Password Last Changed",
    "Last Login",
    "Failed Logins",
    "Reset Link",
    "Notes",
]

_TEMPLATE_HELP = [
    ("Application", "Yes", "The application the account belongs to, e.g. ATTCC."),
    ("App Code", "No", "MOTS / app code, e.g. 31599."),
    ("Type", "Yes", "ITSERVICES, TESTITSERVICES, AAF or JFrog."),
    ("Directory", "No", "AAF: Prod or IST (default Prod). JFrog: Bronze or Gold (default Bronze)."),
    ("Account", "Yes", "Domain: m52142. AAF: m52142@com.att.attcc. JFrog: m52142@att.com."),
    ("Owner", "No", "Who owns the account — name or email."),
    ("Expiry Date", "No", "YYYY-MM-DD or MM/DD/YYYY, or Never."),
    ("Password Last Changed", "No", "YYYY-MM-DD or MM/DD/YYYY."),
    ("Last Login", "No", "YYYY-MM-DD or MM/DD/YYYY."),
    ("Failed Logins", "No", "A whole number."),
    ("Reset Link", "No", "https:// link to where the password is reset."),
    ("Notes", "No", "Anything else: purpose, consumers, rotation steps."),
]


def import_template_xlsx() -> bytes:
    """Empty "Accounts" sheet with dropdowns for Type / Directory, plus a "How to fill" sheet."""
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Accounts"
    for col, header in enumerate(TEMPLATE_HEADERS, start=1):
        ws.cell(row=1, column=col, value=header).font = Font(bold=True)
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = max(len(header) + 4, 16)
    dropdowns = (("C", "ITSERVICES,TESTITSERVICES,AAF,JFrog"), ("D", "ITSERVICES,TESTITSERVICES,Prod,IST,Bronze,Gold"))
    for column, choices in dropdowns:
        validation = DataValidation(type="list", formula1=f'"{choices}"', allow_blank=True)
        ws.add_data_validation(validation)
        validation.add(f"{column}2:{column}{IMPORT_MAX_ROWS + 1}")
    ws.freeze_panes = "A2"

    guide = workbook.create_sheet("How to fill")
    for col, header in enumerate(("Column", "Required", "What to enter"), start=1):
        guide.cell(row=1, column=col, value=header).font = Font(bold=True)
    for row, entry in enumerate(_TEMPLATE_HELP, start=2):
        for col, text in enumerate(entry, start=1):
            guide.cell(row=row, column=col, value=text)
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 10
    guide.column_dimensions["C"].width = 80
    guide.cell(
        row=len(_TEMPLATE_HELP) + 3,
        column=1,
        value=(
            "Never put a password or token in this file. Re-importing updates existing accounts; "
            "blank cells keep what is recorded."
        ),
    ).font = Font(bold=True)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ── Filtering (shared by the grid's export) ────────────────────────────


@dataclass(frozen=True)
class AccountFilter:
    application: str | None = None
    account_type: str | None = None
    directory: str | None = None
    view: str | None = None  # expired | expiring | failing
    q: str | None = None


def is_expired(account: CredentialAccount, today: date) -> bool:
    return bool(account.expiry_date and not account.never_expires and account.expiry_date < today)


def is_expiring(account: CredentialAccount, today: date) -> bool:
    days = _days_left(account, today)
    return days is not None and 0 <= days <= EXPIRING_WINDOW_DAYS


def is_failing(account: CredentialAccount) -> bool:
    return bool(account.failed_logins) or account.last_check_status in FAILING_STATUSES


def filter_accounts(accounts: list[CredentialAccount], flt: AccountFilter, *, today: date) -> list[CredentialAccount]:
    """The grid's filters, so an export holds exactly the rows on screen."""
    query = (flt.q or "").strip().lower()
    result = []
    for account in accounts:
        if flt.application and account.application != flt.application:
            continue
        if flt.account_type and account.account_type != flt.account_type:
            continue
        if flt.directory and account.directory != flt.directory:
            continue
        if flt.view == "expired" and not is_expired(account, today):
            continue
        if flt.view == "expiring" and not is_expiring(account, today):
            continue
        if flt.view == "failing" and not is_failing(account):
            continue
        if query:
            # Field by field, exactly as the grid's search box matches.
            fields = (
                account.application,
                account.app_code,
                account.account_name,
                account.display_name,
                account.owner,
                account.notes,
                f"{TYPE_LABELS[account.account_type]} {DIRECTORY_LABELS.get(account.directory, account.directory)}",
            )
            if not any(query in str(value).lower() for value in fields if value):
                continue
        result.append(account)
    return result


# ── Service ────────────────────────────────────────────────────────────


class CredentialInventoryService:
    def __init__(self, db: AsyncSession) -> None:
        self._db = db

    async def list_accounts(self) -> list[CredentialAccount]:
        result = await self._db.execute(
            select(CredentialAccount).order_by(CredentialAccount.application, CredentialAccount.account_key)
        )
        return list(result.scalars().all())

    async def get(self, account_id: int) -> CredentialAccount:
        account = await self._db.get(CredentialAccount, account_id)
        if account is None:
            raise LookupError(f"Account {account_id} is not in the inventory.")
        return account

    async def _find(self, account_type: str, directory: str, account_name: str) -> CredentialAccount | None:
        result = await self._db.execute(
            select(CredentialAccount).where(
                CredentialAccount.account_type == account_type,
                CredentialAccount.directory == directory,
                CredentialAccount.account_key == account_name.strip().lower(),
            )
        )
        return result.scalar_one_or_none()

    async def create(self, payload: CredentialAccountIn, user: UserContext) -> CredentialAccount:
        if await self._find(payload.account_type, payload.directory, payload.account_name):
            raise DuplicateAccountError(f"{account_label(payload)} is already in the inventory.")
        account = CredentialAccount(created_by=user.email or user.user_id)
        self._apply(account, payload.model_dump(), user, replace=True)
        self._db.add(account)
        await self._db.commit()
        await self._db.refresh(account)
        return account

    async def update(
        self, account_id: int, payload: CredentialAccountIn, user: UserContext
    ) -> tuple[CredentialAccount, list[str]]:
        """Replace every field.  Returns the account and the names of the fields that changed."""
        account = await self.get(account_id)
        clash = await self._find(payload.account_type, payload.directory, payload.account_name)
        if clash is not None and clash.id != account_id:
            raise DuplicateAccountError(f"{account_label(payload)} is already in the inventory.")
        changed = self._apply(account, payload.model_dump(), user, replace=True)
        await self._db.commit()
        await self._db.refresh(account)
        return account, changed

    async def delete(self, account_id: int) -> CredentialAccount:
        account = await self.get(account_id)
        await self._db.delete(account)
        await self._db.commit()
        return account

    @staticmethod
    def _apply(account: CredentialAccount, values: dict[str, Any], user: UserContext, *, replace: bool) -> list[str]:
        """Copy values onto the row; with ``replace=False`` a None keeps what is there."""
        changed = []
        for field, value in values.items():
            if value is None and not replace:
                continue
            if getattr(account, field) != value:
                setattr(account, field, value)
                changed.append(field)
        if values.get("never_expires"):
            account.expiry_date = None
        account.account_key = account.account_name.strip().lower()
        account.updated_by = user.email or user.user_id
        return changed

    # ── Import ─────────────────────────────────────────────────────────

    async def preview_import(self, filename: str, content: bytes) -> ImportPreview:
        """Validate every row and say what importing it would do.  Writes nothing."""
        existing = {(a.account_type, a.directory, a.account_key) for a in await self.list_accounts()}
        seen: dict[tuple[str, str, str], int] = {}
        rows: list[ImportPreviewRow] = []
        for number, raw in read_import_file(filename, content):
            account, errors = parse_import_row(raw)
            label = (
                " / ".join(
                    part for part in (_cell_text(raw.get("application")), _cell_text(raw.get("account_name"))) if part
                )
                or "(blank)"
            )
            if account is None:
                rows.append(ImportPreviewRow(row=number, action="invalid", label=label, errors=errors))
                continue
            key = (account.account_type, account.directory, account.account_name.lower())
            if key in seen:
                rows.append(
                    ImportPreviewRow(
                        row=number,
                        action="invalid",
                        label=label,
                        errors=[f"Same account as row {seen[key]}"],
                    )
                )
                continue
            seen[key] = number
            rows.append(
                ImportPreviewRow(
                    row=number,
                    action="update" if key in existing else "create",
                    account=account,
                    label=account_label(account),
                )
            )
        return ImportPreview(
            rows=rows,
            total=len(rows),
            create_count=sum(1 for r in rows if r.action == "create"),
            update_count=sum(1 for r in rows if r.action == "update"),
            invalid_count=sum(1 for r in rows if r.action == "invalid"),
        )

    async def apply_import(self, rows: list[CredentialAccountImportRow], user: UserContext) -> ImportApplyResult:
        """Create new accounts and fill in existing ones, in one transaction."""
        existing = {(a.account_type, a.directory, a.account_key): a for a in await self.list_accounts()}
        created = updated = 0
        for row in rows:
            key = (row.account_type, row.directory, row.account_name.strip().lower())
            values = row.model_dump()
            account = existing.get(key)
            if account is None:
                account = CredentialAccount(created_by=user.email or user.user_id, never_expires=False)
                self._apply(account, values, user, replace=False)
                self._db.add(account)
                existing[key] = account
                created += 1
            elif self._apply(account, values, user, replace=False):
                updated += 1
        await self._db.commit()
        return ImportApplyResult(created=created, updated=updated)

    # ── Credential Check results ───────────────────────────────────────

    async def record_check(
        self,
        *,
        account_type: str,
        directory: str,
        account_name: str,
        outcome: CheckOutcome,
        checked_by: str,
    ) -> bool:
        """Save a check's verdict and reported details on the listed account.

        True when a row was updated.  An "error" verdict (unreachable host)
        says nothing about the account, so it changes nothing.
        """
        if outcome.status == "error":
            return False
        account = await self._find(account_type, directory, account_name)
        if account is None:
            return False

        account.last_checked_at = datetime.utcnow()
        account.last_check_status = outcome.status
        account.last_checked_by = checked_by
        details = outcome.account
        if details is not None:
            if details.never_expires:
                account.never_expires = True
                account.expiry_date = None
            elif details.expires_at is not None:
                account.never_expires = False
                account.expiry_date = _utc_date(details.expires_at)
            if details.last_used_at is not None:
                account.last_login_date = _utc_date(details.last_used_at)
            if details.last_changed_at is not None:
                account.password_changed_date = _utc_date(details.last_changed_at)
            # AD's displayName names the account; a JFrog "display name" is
            # only the token's description, so it is not kept.
            if account_type == "domain" and details.display_name:
                account.display_name = details.display_name[:255]
        await self._db.commit()
        return True

    # ── Audit ──────────────────────────────────────────────────────────

    async def record_audit(
        self,
        user: UserContext,
        action: str,
        *,
        resource_id: str | None,
        summary: str,
        details: dict | None = None,
        ip_address: str | None = None,
    ) -> None:
        """``audit_logs`` row for a change made from the page.  Never raises."""
        try:
            self._db.add(
                AuditLog(
                    user_id=user.user_id,
                    user_email=user.email,
                    action=action,
                    resource_type=_AUDIT_RESOURCE_TYPE,
                    resource_id=resource_id,
                    details={
                        "page": "CredentialCheckPage",
                        "feature": "credential_inventory",
                        "summary": summary,
                        **(details or {}),
                    },
                    ip_address=ip_address,
                    status="success",
                )
            )
            await self._db.commit()
        except Exception as exc:
            await self._db.rollback()
            logger.warning("credential_inventory_audit_failed", action=action, error=str(exc)[:200])
