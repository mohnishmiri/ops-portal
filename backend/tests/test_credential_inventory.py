"""Credential account inventory: import parsing, export, check write-back, access.

The import is the part most likely to meet messy input — people's mechID
sheets are not written to a template — so most of the coverage is there.
"""

import io
from datetime import UTC, date, datetime

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.auth import get_current_user
from app.core.database import get_db
from app.models.database import AuditLog, CredentialAccount
from app.schemas.auth import UserContext, UserRole
from app.schemas.credential_inventory import CredentialAccountCreate, CredentialAccountImportRow, CredentialAccountIn
from app.schemas.tools import CredentialAccountDetails
from app.services import credential_inventory_service as inv
from app.services.credential_check_service import CheckOutcome, DirectoryLookup

TODAY = date(2026, 10, 12)


def _writer() -> UserContext:
    return UserContext(
        user_id="test-writer",
        object_id="22222222-2222-2222-2222-222222222222",
        display_name="Test Writer",
        email="writer@example.com",
        roles=[UserRole.WRITE],
        raw_roles=["write"],
        tenant_id="tenant-test",
        allowed_subscriptions=[],
    )


def _account(**overrides) -> CredentialAccountIn:
    values = {
        "application": "ATTCC",
        "account_type": "domain",
        "directory": "ITSERVICES",
        "account_name": "m52142",
    }
    values.update(overrides)
    return CredentialAccountIn(**values)


def _xlsx(rows: list[list], header_at: int = 1) -> bytes:
    workbook = Workbook()
    ws = workbook.active
    for _ in range(header_at - 1):
        ws.append(["AT&T Proprietary (Restricted)"])
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ── Reading a row ──────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("type_cell", "directory_cell", "expected"),
    [
        ("ITSERVICES", None, ("domain", "ITSERVICES")),
        ("testitservices", None, ("domain", "TESTITSERVICES")),
        ("Domain", "TESTITSERVICES", ("domain", "TESTITSERVICES")),
        ("AAF", None, ("aaf", "prod")),
        ("AAF", "IST", ("aaf", "nonprod")),
        ("AAF - Non-Prod", None, ("aaf", "nonprod")),
        ("JFrog", None, ("jfrog", "bronze")),
        ("JFrog Gold", None, ("jfrog", "gold")),
        ("Artifactory", "artifact-gold.it.att.com", ("jfrog", "gold")),
        (None, "Gold", ("jfrog", "gold")),  # no Type column: the directory says it
        (None, "ITSERVICES", ("domain", "ITSERVICES")),
    ],
)
def test_import_reads_the_ways_people_write_the_type(type_cell, directory_cell, expected):
    row, errors = inv.parse_import_row(
        {"application": "ATTCC", "account_type": type_cell, "directory": directory_cell, "account_name": "m1"}
    )
    assert errors == []
    assert (row.account_type, row.directory) == expected


@pytest.mark.parametrize(
    ("type_cell", "directory_cell", "fragment"),
    [
        ("Kerberos", None, "is not one of"),
        ("JFrog", "IST", "does not fit a JFrog account"),
        ("AAF", "Platinum", "is not recognised"),
    ],
)
def test_import_names_a_bad_type_or_directory(type_cell, directory_cell, fragment):
    row, errors = inv.parse_import_row(
        {"application": "ATTCC", "account_type": type_cell, "directory": directory_cell, "account_name": "m1"}
    )
    assert row is None
    assert fragment in errors[0]


def test_import_ignores_columns_the_directory_answers_for():
    """An Expiry Date or Owner column in someone's sheet must not beat AD."""
    row, errors = inv.parse_import_row(
        {
            "application": "ATTCC",
            "account_type": "AAF",
            "account_name": "m1@com.att.attcc",
            "expiry_date": "next tuesday",
            "owner": "Jane Doe",
            "failed_logins": "lots",
        }
    )
    assert errors == []
    assert not hasattr(row, "expiry_date")
    assert not hasattr(row, "owner")


def test_import_collects_every_problem_in_a_row():
    row, errors = inv.parse_import_row({"account_type": "Kerberos", "account_name": "", "reset_url": "ftp://x/y"})
    assert row is None
    assert "Application: required" in errors
    assert "Account: required" in errors
    assert any("is not one of" in e for e in errors)


def test_a_reset_link_must_be_a_web_link():
    """It is rendered as a link — javascript: would run in the reader's browser."""
    row, errors = inv.parse_import_row(
        {
            "application": "ATTCC",
            "account_type": "ITSERVICES",
            "account_name": "m1",
            "reset_url": "javascript:alert(1)",
        }
    )
    assert row is None
    assert errors == ["Reset Link: Reset link must be an http:// or https:// URL"]


# ── Reading a file ─────────────────────────────────────────────────────


def test_import_finds_the_header_below_a_title_so_an_export_imports_back():
    content = _xlsx(
        [["Application", "Type", "Account", "Notes"], ["ATTCC", "ITSERVICES", "m52142", "nightly build"]],
        header_at=4,
    )
    rows = inv.read_import_file("accounts.xlsx", content)
    expected = {
        "application": "ATTCC",
        "account_type": "ITSERVICES",
        "account_name": "m52142",
        "notes": "nightly build",
    }
    assert rows == [(5, expected)]


def test_import_reads_header_aliases():
    content = _xlsx(
        [
            ["App Name", "MOTS ID", "Account Type", "MechID", "Expiration Date", "Comments"],
            ["DWS", 31599, "AAF", "m1", None, "x"],
        ]
    )
    ((_, raw),) = inv.read_import_file("ids.xlsx", content)
    # "Expiration Date" is not a column the import reads — the directory answers for it.
    assert set(raw) == {"application", "app_code", "account_type", "account_name", "notes"}


def test_import_reads_a_windows_excel_csv():
    """Excel on Windows saves "CSV" in the ANSI code page, not UTF-8."""
    content = "Application,Type,Account,Notes\r\nATTCC,ITSERVICES,m1,Café build\r\n".encode("cp1252")
    ((_, raw),) = inv.read_import_file("ids.csv", content)
    assert raw["notes"] == "Café build"


@pytest.mark.parametrize(
    ("filename", "content", "fragment"),
    [
        ("ids.xls", b"\xd0\xcf", "save it as .xlsx or .csv"),
        ("ids.txt", b"x", "Upload an Excel"),
        ("ids.xlsx", b"not a zip", "could not be read"),
        ("ids.csv", b"Name,Value\r\na,b\r\n", "No header row found"),
    ],
)
def test_import_refuses_files_it_cannot_read(filename, content, fragment):
    with pytest.raises(ValueError) as excinfo:
        inv.read_import_file(filename, content)
    assert fragment in str(excinfo.value)


def test_import_caps_the_row_count(monkeypatch):
    monkeypatch.setattr(inv, "IMPORT_MAX_ROWS", 2)
    content = "Application,Type,Account\r\n" + "".join(f"A,ITSERVICES,m{i}\r\n" for i in range(3))
    with pytest.raises(ValueError, match="more than 2 accounts"):
        inv.read_import_file("ids.csv", content.encode())


def test_the_template_imports_cleanly_and_holds_no_example_rows():
    content = inv.import_template_xlsx()
    workbook = load_workbook(io.BytesIO(content))
    assert workbook.sheetnames == ["Accounts", "How to fill"]
    assert inv.read_import_file("template.xlsx", content) == []


# ── Export ─────────────────────────────────────────────────────────────


def _row(**overrides) -> CredentialAccount:
    values = {
        "id": 1,
        "application": "ATTCC",
        "app_code": "31599",
        "account_type": "domain",
        "directory": "ITSERVICES",
        "account_name": "m52142",
        "account_key": "m52142",
        "owner": "m52142 - KAREN RIDGEWAY (kr873u)",
        "expiry_date": date(2026, 12, 14),
        "never_expires": False,
        "failed_logins": 0,
        "notes": None,
        "reset_url": None,
        "last_check_status": None,
        "last_checked_at": None,
        "password_changed_date": None,
        "last_login_date": None,
    }
    values.update(overrides)
    return CredentialAccount(**values)


def test_csv_export_neutralises_formulas():
    content = inv.export_csv([_row(notes='=HYPERLINK("http://x")', owner="+1 call")], today=TODAY).decode("utf-8-sig")
    assert "'=HYPERLINK" in content
    assert "'+1 call" in content


def test_xlsx_export_keeps_a_formula_like_note_as_text():
    content = inv.export_xlsx([_row(notes="=1+1")], today=TODAY, generated_at=datetime(2026, 10, 12, tzinfo=UTC))
    ws = load_workbook(io.BytesIO(content)).active
    notes_col = [c.value for c in ws[4]].index("Notes") + 1
    cell = ws.cell(row=5, column=notes_col)
    assert cell.data_type == "s"
    assert cell.value == "=1+1"


def test_an_export_lists_the_directorys_answers_and_still_imports_back():
    rows = [
        _row(reset_url="https://reset.example.com"),
        _row(id=2, account_type="jfrog", directory="gold", account_name="svc@att.com", never_expires=True),
    ]
    content = inv.export_xlsx(rows, today=TODAY, generated_at=datetime(2026, 10, 12, tzinfo=UTC))

    headers = [c.value for c in load_workbook(io.BytesIO(content)).active[4]]
    assert "Name (from directory)" in headers and "Expiry Date" in headers

    parsed = [inv.parse_import_row(raw) for _, raw in inv.read_import_file("export.xlsx", content)]
    assert all(errors == [] for _, errors in parsed)
    first, second = (row for row, _ in parsed)
    assert (first.account_type, first.directory, first.reset_url) == (
        "domain",
        "ITSERVICES",
        "https://reset.example.com",
    )
    assert (second.account_type, second.directory) == ("jfrog", "gold")


def test_filters_match_the_grid_tiles():
    expired = _row(id=1, expiry_date=date(2026, 10, 1))
    expiring = _row(id=2, expiry_date=date(2026, 10, 30))
    later = _row(id=3, expiry_date=date(2027, 6, 1))
    never = _row(id=4, expiry_date=None, never_expires=True)
    locked = _row(id=5, expiry_date=None, last_check_status="locked")
    failed = _row(id=6, expiry_date=None, failed_logins=3, last_check_status="valid")
    everything = [expired, expiring, later, never, locked, failed]

    def ids(view):
        return [a.id for a in inv.filter_accounts(everything, inv.AccountFilter(view=view), today=TODAY)]

    assert ids("expired") == [1]
    assert ids("expiring") == [2]
    assert ids("failing") == [5, 6]
    # Search reaches the name the directory gave, not just what was typed in.
    searched = inv.filter_accounts(everything, inv.AccountFilter(q="ridgeway"), today=TODAY)
    assert [a.id for a in searched] == [1, 2, 3, 4, 5, 6]


# ── Service: create, import, check write-back ──────────────────────────


@pytest.mark.anyio
async def test_the_same_account_cannot_be_listed_twice_in_any_case(db_session):
    service = inv.CredentialInventoryService(db_session)
    await service.create(_account(), _writer())
    with pytest.raises(inv.DuplicateAccountError):
        await service.create(_account(account_name="M52142", application="DWS"), _writer())
    # Same name in another directory is another account.
    await service.create(_account(directory="TESTITSERVICES"), _writer())


@pytest.mark.anyio
async def test_preview_says_create_update_or_why_not(db_session):
    service = inv.CredentialInventoryService(db_session)
    await service.create(_account(), _writer())
    content = (
        b"Application,Type,Account,Failed Logins\r\n"
        b"ATTCC,ITSERVICES,M52142,2\r\n"
        b"ATTCC,AAF,m1@com.att.attcc,\r\n"
        b"ATTCC,aaf,M1@com.att.attcc,\r\n"
        b",JFrog,svc,\r\n"
    )

    preview = await service.preview_import("ids.csv", content)

    assert [(r.row, r.action) for r in preview.rows] == [(2, "update"), (3, "create"), (4, "invalid"), (5, "invalid")]
    assert preview.rows[2].errors == ["Same account as row 3"]
    assert preview.rows[3].errors == ["Application: required"]
    assert (preview.create_count, preview.update_count, preview.invalid_count) == (1, 1, 2)


@pytest.mark.anyio
async def test_an_import_fills_in_but_never_erases_or_touches_the_directorys_answers(db_session):
    service = inv.CredentialInventoryService(db_session)
    created = await service.create(
        CredentialAccountCreate(
            **_account(notes="rotate via CyberArk").model_dump(),
            owner="KAREN RIDGEWAY",
            expiry_date=date(2026, 12, 14),
            failed_logins=1,
        ),
        _writer(),
    )

    result = await service.apply_import(
        [
            CredentialAccountImportRow(
                application="Commissions",
                account_type="domain",
                directory="ITSERVICES",
                account_name="M52142",
            ),
            CredentialAccountImportRow(application="DWS", account_type="aaf", directory="prod", account_name="m1"),
        ],
        _writer(),
    )

    assert (result.created, result.updated) == (1, 1)
    row = await service.get(created.id)
    assert row.application == "Commissions"
    assert row.notes == "rotate via CyberArk"  # blank in the file: kept
    # Nothing in a spreadsheet can overwrite what the directory said.
    assert (row.owner, row.expiry_date, row.failed_logins) == ("KAREN RIDGEWAY", date(2026, 12, 14), 1)


@pytest.mark.anyio
async def test_a_check_writes_its_verdict_and_details_onto_the_listed_account(db_session):
    service = inv.CredentialInventoryService(db_session)
    created = await service.create(_account(), _writer())
    outcome = CheckOutcome(
        valid=True,
        status="valid",
        message="ok",
        account=CredentialAccountDetails(
            display_name="m52142 - KAREN RIDGEWAY (kr873u)",
            expires_at=datetime(2026, 12, 14, 15, 0, tzinfo=UTC),
            last_used_at=datetime(2026, 10, 1, 8, 0, tzinfo=UTC),
            last_changed_at=datetime(2026, 9, 15, tzinfo=UTC),
            failed_sign_ins=2,
        ),
    )

    updated = await service.record_check(
        account_type="domain", directory="ITSERVICES", account_name="M52142", outcome=outcome, checked_by="w@x.com"
    )

    assert updated is True
    row = await service.get(created.id)
    assert (row.last_check_status, row.last_checked_by) == ("valid", "w@x.com")
    assert row.expiry_date == date(2026, 12, 14)
    assert row.last_login_date == date(2026, 10, 1)
    assert row.password_changed_date == date(2026, 9, 15)
    assert row.failed_logins == 2
    # The directory's name for the account is what the grid shows as its owner.
    assert row.owner == "m52142 - KAREN RIDGEWAY (kr873u)"


@pytest.mark.anyio
async def test_an_unreachable_check_and_an_unlisted_account_change_nothing(db_session):
    service = inv.CredentialInventoryService(db_session)
    created = await service.create(_account(), _writer())
    error = CheckOutcome(valid=False, status="error", message="timed out")

    assert (
        await service.record_check(
            account_type="domain", directory="ITSERVICES", account_name="m52142", outcome=error, checked_by="w"
        )
        is False
    )
    assert (
        await service.record_check(
            account_type="domain",
            directory="ITSERVICES",
            account_name="someone-else",
            outcome=CheckOutcome(valid=True, status="valid", message="ok"),
            checked_by="w",
        )
        is False
    )
    assert (await service.get(created.id)).last_check_status is None


# ── API ────────────────────────────────────────────────────────────────


@pytest.fixture
async def writer_client(app, db_session):
    async def _db():
        yield db_session

    async def _user():
        return _writer()

    app.dependency_overrides[get_db] = _db
    app.dependency_overrides[get_current_user] = _user
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        yield ac
    app.dependency_overrides.clear()


_BODY = {"application": "ATTCC", "account_type": "domain", "directory": "itservices", "account_name": "m52142"}


@pytest.mark.anyio
async def test_a_reader_sees_the_inventory_but_cannot_change_it(reader_client):
    assert (await reader_client.get("/api/v1/tools/accounts")).status_code == 200
    assert (await reader_client.post("/api/v1/tools/accounts", json=_BODY)).status_code == 403
    files = {"file": ("ids.csv", b"Application,Type,Account\r\nA,ITSERVICES,m1\r\n", "text/csv")}
    assert (await reader_client.post("/api/v1/tools/accounts/import/preview", files=files)).status_code == 403


@pytest.mark.anyio
async def test_adding_an_account_is_audited_and_never_takes_a_password(writer_client, db_session):
    response = await writer_client.post("/api/v1/tools/accounts", json={**_BODY, "password": "hunter2"})

    assert response.status_code == 201
    assert response.json()["directory"] == "ITSERVICES"
    assert "password" not in response.json()
    audit = (await db_session.execute(select(AuditLog))).scalars().one()
    assert audit.action == "create_credential_account"
    assert audit.details["feature"] == "credential_inventory"
    assert "hunter2" not in str(audit.details)

    duplicate = await writer_client.post("/api/v1/tools/accounts", json={**_BODY, "account_name": "M52142"})
    assert duplicate.status_code == 409


@pytest.mark.anyio
async def test_export_returns_the_filtered_rows(writer_client):
    await writer_client.post("/api/v1/tools/accounts", json=_BODY)
    await writer_client.post(
        "/api/v1/tools/accounts", json={**_BODY, "application": "DWS", "account_type": "jfrog", "directory": "gold"}
    )

    response = await writer_client.get("/api/v1/tools/accounts/export", params={"format": "csv", "application": "DWS"})

    assert response.status_code == 200
    assert "attachment" in response.headers["content-disposition"]
    lines = response.content.decode("utf-8-sig").strip().splitlines()
    assert len(lines) == 2 and "Gold" in lines[1]


@pytest.mark.anyio
async def test_a_check_of_a_listed_account_updates_its_row(writer_client, monkeypatch):
    from app.api.v1.endpoints import tools

    await writer_client.post(
        "/api/v1/tools/accounts",
        json={**_BODY, "account_type": "jfrog", "directory": "gold", "account_name": "svc@att.com"},
    )
    seen = {}

    async def _fake_check(username, token, instance="bronze"):
        seen["instance"] = instance
        return CheckOutcome(valid=True, status="valid", message="ok")

    monkeypatch.setattr(tools._service, "check_jfrog", _fake_check)

    response = await writer_client.post(
        "/api/v1/tools/credential-checks/jfrog",
        json={"instance": "gold", "username": "svc@att.com", "token": "t"},
    )

    assert response.status_code == 200
    assert response.json()["inventory_updated"] is True
    assert response.json()["target"] == "svc@att.com (Gold)"
    assert seen["instance"] == "gold"
    rows = (await writer_client.get("/api/v1/tools/accounts")).json()
    assert rows[0]["last_check_status"] == "valid"


# ── The directory's name for the account ──────────────────────────────


@pytest.mark.anyio
async def test_only_a_domain_check_names_the_owner(db_session):
    """AD's displayName names the person behind a mechID; a JFrog one is a token description."""
    service = inv.CredentialInventoryService(db_session)
    domain = await service.create(_account(), _writer())
    jfrog = await service.create(_account(account_type="jfrog", directory="bronze", account_name="svc"), _writer())

    async def _check(account_type, directory, name, display):
        await service.record_check(
            account_type=account_type,
            directory=directory,
            account_name=name,
            outcome=CheckOutcome(
                valid=True, status="valid", message="ok", account=CredentialAccountDetails(display_name=display)
            ),
            checked_by="w",
        )

    await _check("domain", "ITSERVICES", "m52142", "m52142 - KAREN RIDGEWAY (kr873u)")
    await _check("jfrog", "bronze", "svc", "pipeline token")  # a token description, not a name

    assert (await service.get(domain.id)).owner == "m52142 - KAREN RIDGEWAY (kr873u)"
    assert (await service.get(jfrog.id)).owner is None


@pytest.mark.anyio
async def test_editing_an_account_cannot_touch_what_the_directory_said(db_session):
    """The form has no field for these, and the API ignores one if it is sent anyway."""
    service = inv.CredentialInventoryService(db_session)
    created = await service.create(
        CredentialAccountCreate(
            **_account().model_dump(), owner="KAREN RIDGEWAY", expiry_date=date(2027, 9, 30), failed_logins=3
        ),
        _writer(),
    )

    await service.update(
        created.id,
        CredentialAccountIn(**_account(notes="moved to the new pipeline").model_dump()),
        _writer(),
    )

    row = await service.get(created.id)
    assert row.notes == "moved to the new pipeline"
    assert (row.owner, row.expiry_date, row.failed_logins) == ("KAREN RIDGEWAY", date(2027, 9, 30), 3)


@pytest.mark.anyio
async def test_a_check_says_what_it_checked_so_the_page_can_list_it(writer_client, monkeypatch):
    from app.api.v1.endpoints import tools

    async def _fake_check(domain, username, password):
        return CheckOutcome(valid=True, status="valid", message="ok")

    monkeypatch.setattr(tools._service, "check_domain", _fake_check)

    response = await writer_client.post(
        "/api/v1/tools/credential-checks/domain",
        json={"domain": "TESTITSERVICES", "username": "m47355", "password": "p"},
    )

    body = response.json()
    assert (body["directory"], body["account_name"], body["inventory_updated"]) == ("TESTITSERVICES", "m47355", False)


@pytest.mark.anyio
async def test_adding_from_a_check_result_keeps_its_details(writer_client):
    response = await writer_client.post(
        "/api/v1/tools/accounts",
        json={
            **_BODY,
            "expiry_date": "2027-09-30",
            "password_changed_date": "2026-09-30",
            "last_login_date": "2026-10-11",
            "owner": "m52142 - KAREN RIDGEWAY (kr873u)",
            "failed_logins": 0,
        },
    )

    assert response.status_code == 201
    body = response.json()
    assert (body["expiry_date"], body["owner"]) == ("2027-09-30", "m52142 - KAREN RIDGEWAY (kr873u)")
    assert (body["last_login_date"], body["failed_logins"]) == ("2026-10-11", 0)


# ── Filling a row in from the directory ────────────────────────────────


class _FakeChecks:
    """Stands in for the directory: records what was asked, answers as told."""

    def __init__(self, lookup: DirectoryLookup):
        self.lookup = lookup
        self.asked: list[tuple[str, str]] = []

    async def lookup_domain_account(self, domain: str, username: str) -> DirectoryLookup:
        self.asked.append((domain, username))
        return self.lookup


def _found(**overrides) -> DirectoryLookup:
    details = CredentialAccountDetails(
        display_name="m47355 - KAREN RIDGEWAY (kr873u)",
        expires_at=datetime(2027, 9, 30, 15, 0, tzinfo=UTC),
        last_changed_at=datetime(2026, 9, 30, tzinfo=UTC),
        last_used_at=datetime(2026, 10, 11, 8, 0, tzinfo=UTC),
        failed_sign_ins=0,
        **overrides,
    )
    return DirectoryLookup("ok", None, details)


@pytest.mark.anyio
async def test_a_domain_account_is_filled_in_from_the_directory(db_session):
    checks = _FakeChecks(_found())
    service = inv.CredentialInventoryService(db_session, checks)
    created = await service.create(_account(account_name="m47355"), _writer())

    lookup = await service.sync_from_directory(created)

    assert lookup.status == "ok"
    assert checks.asked == [("ITSERVICES", "m47355")]
    row = await service.get(created.id)
    assert row.owner == "m47355 - KAREN RIDGEWAY (kr873u)"
    assert row.expiry_date == date(2027, 9, 30)
    assert row.password_changed_date == date(2026, 9, 30)
    assert row.last_login_date == date(2026, 10, 11)
    assert row.failed_logins == 0
    assert row.directory_synced_at is not None
    # Reading the account proves nothing about its password.
    assert (row.last_check_status, row.last_checked_at) == (None, None)


@pytest.mark.anyio
@pytest.mark.parametrize("lookup", [DirectoryLookup("unconfigured", "no reader"), DirectoryLookup("error", "down")])
async def test_a_directory_that_cannot_answer_leaves_the_row_alone(db_session, lookup):
    service = inv.CredentialInventoryService(db_session, _FakeChecks(lookup))
    created = await service.create(_account(), _writer())

    result = await service.sync_from_directory(created)

    assert result.status == lookup.status
    row = await service.get(created.id)
    assert (row.owner, row.expiry_date, row.directory_synced_at) == (None, None, None)


@pytest.mark.anyio
async def test_aaf_and_jfrog_accounts_say_a_check_is_the_way_in(db_session):
    checks = _FakeChecks(_found())
    service = inv.CredentialInventoryService(db_session, checks)
    jfrog = await service.create(
        _account(account_type="jfrog", directory="gold", account_name="svc@att.com"), _writer()
    )

    lookup = await service.sync_from_directory(jfrog)

    assert lookup.status == "unsupported"
    assert "check of the account itself" in lookup.reason
    assert checks.asked == []  # no directory was asked


@pytest.mark.anyio
async def test_adding_an_itservices_account_fills_the_grid_in_at_once(writer_client, monkeypatch):
    """What the user asked for: list the account, and its details are there."""
    from app.api.v1.endpoints import credential_accounts as endpoint

    async def _sync(self, account):
        inv._apply_directory_details(account, _found().details, naming=True)
        account.directory_synced_at = datetime.utcnow()
        return _found()

    monkeypatch.setattr(endpoint.CredentialInventoryService, "sync_from_directory", _sync)

    response = await writer_client.post("/api/v1/tools/accounts", json={**_BODY, "account_name": "m47355"})

    assert response.status_code == 201
    body = response.json()
    assert body["owner"] == "m47355 - KAREN RIDGEWAY (kr873u)"
    assert body["expiry_date"] == "2027-09-30"
    assert body["directory_synced_at"] is not None


@pytest.mark.anyio
async def test_adding_from_a_check_result_does_not_read_the_directory_again(writer_client, monkeypatch):
    """The check that just ran already read it — a second read would be wasted."""
    from app.api.v1.endpoints import credential_accounts as endpoint

    called = []

    async def _sync(self, account):
        called.append(account)
        return DirectoryLookup("ok")

    monkeypatch.setattr(endpoint.CredentialInventoryService, "sync_from_directory", _sync)

    await writer_client.post(
        "/api/v1/tools/accounts", json={**_BODY, "account_name": "m47355", "expiry_date": "2027-09-30"}
    )

    assert called == []


@pytest.mark.anyio
async def test_refresh_reports_why_the_directory_could_not_answer(writer_client, monkeypatch):
    from app.api.v1.endpoints import credential_accounts as endpoint

    created = await writer_client.post("/api/v1/tools/accounts", json=_BODY)

    async def _sync(self, account):
        return DirectoryLookup("unconfigured", "No directory reader is configured for ITSERVICES.")

    monkeypatch.setattr(endpoint.CredentialInventoryService, "sync_from_directory", _sync)

    response = await writer_client.post(f"/api/v1/tools/accounts/{created.json()['id']}/refresh")

    assert response.status_code == 200
    assert response.json()["status"] == "unconfigured"
    assert "No directory reader" in response.json()["reason"]


@pytest.mark.anyio
async def test_a_reader_cannot_refresh(reader_client):
    assert (await reader_client.post("/api/v1/tools/accounts/1/refresh")).status_code == 403
