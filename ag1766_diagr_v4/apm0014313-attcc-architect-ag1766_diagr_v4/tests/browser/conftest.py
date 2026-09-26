"""
Browser-suite harness.

Runs the real application in a subprocess against a throwaway database and
drives it with Playwright. Seeding goes through the service layer, never the
UI, except in the journey that exists to test creation.

Landmines this harness exists to neutralise — each one was hit for real
while testing this application by hand:

1. ``persistence/migrations/env.py`` calls ``load_dotenv``, so an *unset*
   ``DATABASE_URL`` is silently backfilled from the repository ``.env`` and
   every "isolated" test then writes to the developer's ``local.db``. The
   server subprocess is therefore given an explicit ``DATABASE_URL``, which
   takes precedence over the dotenv fallback.
2. A catalog with no published options renders every controlled editor as an
   empty ``<select>``, so answers cannot be displayed or selected. The
   catalog fixture asserts options exist rather than assuming it.
3. Evidence with no matching ``CORRELATION`` identifier quarantines every
   import, suppressing all candidates. The application fixture seeds one.
4. Evidence is deduplicated by content hash and imports are idempotent per
   evidence item, so re-using identical bytes silently returns an earlier
   run instead of importing. Every generated workbook embeds a unique token.
5. Optimistic concurrency compares the *revision number*, so tokens must be
   read from the rendered form and never hardcoded.
"""

from __future__ import annotations

import os
import shutil
import socket
import subprocess
import sys
import time
import uuid
from collections.abc import Iterator  # noqa: TC003
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

import httpx
import pytest
from openpyxl import Workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
PYTHON = sys.executable
STARTUP_TIMEOUT_SECONDS = 60


def _free_port() -> int:
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return int(sock.getsockname()[1])


class RestartableAppServer:
    """A disposable real server that can restart against its same SQLite/store."""

    def __init__(
        self,
        environment: dict[str, str],
        working_dir: Path,
        log_dir: Path,
    ) -> None:
        self._environment = environment
        self._working_dir = working_dir
        self._log_dir = log_dir
        self._port = _free_port()
        self._process: subprocess.Popen[bytes] | None = None
        self._log_file: Any | None = None
        self._restart_count = 0

    @property
    def base_url(self) -> str:
        return f"http://127.0.0.1:{self._port}"

    def start(self) -> None:
        if self._process is not None:
            raise RuntimeError("restartable server is already running")
        log_path = self._log_dir / f"server-{self._restart_count}.log"
        self._log_file = log_path.open("wb")
        self._process = subprocess.Popen(  # noqa: S603
            [
                PYTHON,
                "-m",
                "uvicorn",
                "migration_intake.main:get_app",
                "--factory",
                "--host",
                "127.0.0.1",
                "--port",
                str(self._port),
            ],
            cwd=self._working_dir,
            env=self._environment,
            stdout=self._log_file,
            stderr=subprocess.STDOUT,
        )
        deadline = time.monotonic() + STARTUP_TIMEOUT_SECONDS
        while time.monotonic() < deadline:
            if self._process.poll() is not None:
                self._close_log()
                raise RuntimeError(
                    "restartable server exited during startup:\n"
                    f"{log_path.read_text(errors='replace')}"
                )
            try:
                with socket.create_connection(("127.0.0.1", self._port), timeout=2.0):
                    return
            except OSError:
                time.sleep(0.3)
        self.stop()
        raise RuntimeError("restartable server did not become ready in time")

    def restart(self) -> None:
        self.stop()
        self._restart_count += 1
        self.start()

    def stop(self) -> None:
        if self._process is None:
            return
        self._process.terminate()
        try:
            self._process.wait(timeout=10)
        except subprocess.TimeoutExpired:
            self._process.kill()
        self._process = None
        self._close_log()

    def _close_log(self) -> None:
        if self._log_file is not None:
            self._log_file.close()
            self._log_file = None


@pytest.fixture(scope="session")
def browser_db(tmp_path_factory: pytest.TempPathFactory) -> str:
    """Return a SQLite URL for a throwaway database used only by this suite."""
    directory = tmp_path_factory.mktemp("browser-db")
    return f"sqlite:///{(directory / 'browser.db').as_posix()}"


@pytest.fixture(scope="session")
def evidence_root(tmp_path_factory: pytest.TempPathFactory) -> Path:
    return tmp_path_factory.mktemp("browser-evidence")


@pytest.fixture(scope="session")
def server_working_dir(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Keep server startup away from the repository's developer-only .env."""
    directory = tmp_path_factory.mktemp("browser-server-cwd")
    shutil.copytree(
        REPO_ROOT / "src" / "migration_intake" / "web" / "static",
        directory / "static",
    )
    return directory


@pytest.fixture(scope="session")
def server_env(
    browser_db: str,
    evidence_root: Path,
    server_working_dir: Path,
) -> dict[str, str]:
    """
    Environment for the app under test.

    ``DATABASE_URL`` is explicit on purpose: leaving it unset lets Alembic's
    ``load_dotenv`` fallback point the suite at the developer's local
    database (see module docstring, landmine 1).
    """
    env = dict(os.environ)
    for key in (
        "AWS_OUTPOST_DATABASE_URL",
        "DATABASE_URL",
        "AWS_OUTPOST_EVIDENCE_ROOT",
        "EVIDENCE_ROOT",
        "AWS_OUTPOST_STATIC_ASSETS_DIR",
        "STATIC_ASSETS_DIR",
        "AWS_OUTPOST_APP_ENV",
        "APP_ENV",
        "AWS_OUTPOST_ACTOR_ID",
        "ACTOR_ID",
        "AWS_OUTPOST_ACTOR_DISPLAY_NAME",
        "ACTOR_DISPLAY_NAME",
        "AWS_OUTPOST_CSRF_SECRET",
        "CSRF_SECRET",
        "AWS_OUTPOST_ORACLE_CLIENT_LIB_DIR",
        "ORACLE_CLIENT_LIB_DIR",
        "AWS_OUTPOST_TOPOLOGY_GENERATION_ENABLED",
        "TOPOLOGY_GENERATION_ENABLED",
        "AWS_OUTPOST_LLM_ENABLED",
        "LLM_ENABLED",
        "AWS_OUTPOST_LLM_OUTBOUND_ENABLED",
        "LLM_OUTBOUND_ENABLED",
        "AWS_OUTPOST_LLM_PROVIDER",
        "CANDIDATE_MAPPER_PROVIDER",
        "AWS_OUTPOST_LLM_PROFILE",
        "LLM_PROFILE",
    ):
        env.pop(key, None)
    source_paths = [str(REPO_ROOT / "src"), str(REPO_ROOT)]
    if inherited_pythonpath := env.get("PYTHONPATH"):
        source_paths.append(inherited_pythonpath)
    env.update(
        {
            "DATABASE_URL": browser_db,
            "AWS_OUTPOST_DATABASE_URL": browser_db,
            "EVIDENCE_ROOT": str(evidence_root),
            "AWS_OUTPOST_EVIDENCE_ROOT": str(evidence_root),
            "AWS_OUTPOST_STATIC_ASSETS_DIR": str(server_working_dir / "static"),
            "APP_ENV": "test",
            "AWS_OUTPOST_APP_ENV": "test",
            "CSRF_SECRET": "browser-suite-csrf-secret",
            "AWS_OUTPOST_CSRF_SECRET": "browser-suite-csrf-secret",
            "ACTOR_ID": str(uuid.uuid4()),
            "AWS_OUTPOST_ACTOR_ID": str(uuid.uuid4()),
            "ACTOR_DISPLAY_NAME": "Browser Suite Actor",
            "AWS_OUTPOST_ACTOR_DISPLAY_NAME": "Browser Suite Actor",
            "AWS_OUTPOST_LLM_ENABLED": "false",
            "LLM_ENABLED": "false",
            "AWS_OUTPOST_LLM_OUTBOUND_ENABLED": "false",
            "LLM_OUTBOUND_ENABLED": "false",
            "AWS_OUTPOST_LLM_PROVIDER": "mock",
            "CANDIDATE_MAPPER_PROVIDER": "mock",
            "AWS_OUTPOST_LLM_PROFILE": "disabled",
            "LLM_PROFILE": "disabled",
            "AWS_OUTPOST_TOPOLOGY_GENERATION_ENABLED": "false",
            "TOPOLOGY_GENERATION_ENABLED": "false",
            "PYTHONPATH": os.pathsep.join(source_paths),
            "PYTHONUNBUFFERED": "1",
            "LOG_LEVEL": "WARNING",
            # Bypass corporate proxy for localhost connections
            "NO_PROXY": "localhost,127.0.0.1,::1",
            "no_proxy": "localhost,127.0.0.1,::1",
        }
    )
    # Remove proxy settings that could interfere with localhost connections
    for proxy_key in ("HTTP_PROXY", "HTTPS_PROXY", "http_proxy", "https_proxy"):
        env.pop(proxy_key, None)
    # Both aliases resolve through Settings; use one configured actor for all
    # server and seed operations rather than two independently generated IDs.
    env["AWS_OUTPOST_ACTOR_ID"] = env["ACTOR_ID"]
    return env


@pytest.fixture(scope="session")
def migrated_db(server_env: dict[str, str], server_working_dir: Path) -> str:
    """Create the schema and publish the packaged catalog with its options."""
    subprocess.run(  # noqa: S603
        [PYTHON, "-m", "alembic", "upgrade", "head"],
        # alembic.ini's script_location is repository-relative. This migration
        # process does not construct the application or Oracle client; explicit
        # database aliases still override dotenv for its migration boundary.
        cwd=REPO_ROOT,
        env=server_env,
        check=True,
        capture_output=True,
    )
    subprocess.run(  # noqa: S603
        [PYTHON, "-m", "migration_intake.catalog.bootstrap"],
        cwd=server_working_dir,
        env=server_env,
        check=True,
        capture_output=True,
    )

    from sqlalchemy import text

    from migration_intake.persistence.database import create_engine_from_url

    engine = create_engine_from_url(server_env["DATABASE_URL"])
    with engine.connect() as connection:
        questions = connection.execute(text("SELECT count(*) FROM cat_questions")).scalar()
        options = connection.execute(text("SELECT count(*) FROM cat_options")).scalar()
    engine.dispose()

    assert questions, "catalog published no questions"
    # Landmine 2: without options every controlled editor renders empty, so a
    # green suite would prove nothing about answer display.
    assert options, "catalog published no allowed values (cat_options is empty)"
    return server_env["DATABASE_URL"]


@pytest.fixture(scope="session")
def app_server(
    server_env: dict[str, str],
    migrated_db: str,
    server_working_dir: Path,
    tmp_path_factory: pytest.TempPathFactory,
) -> Iterator[str]:
    """
    Start the real app on a free port and yield its base URL.

    Server output goes to a file rather than a pipe: uvicorn logs every
    request, and an undrained ``subprocess.PIPE`` fills its buffer after a
    few dozen requests and blocks the server mid-suite.
    """
    # Bypass corporate proxy for localhost connections
    os.environ["NO_PROXY"] = "localhost,127.0.0.1,::1"
    os.environ["no_proxy"] = "localhost,127.0.0.1,::1"

    port = _free_port()
    log_path = tmp_path_factory.mktemp("browser-server") / "server.log"
    log_file = log_path.open("wb")
    process = subprocess.Popen(  # noqa: S603
        [
            PYTHON,
            "-m",
            "uvicorn",
            "migration_intake.main:get_app",
            "--factory",
            "--host",
            "127.0.0.1",
            "--port",
            str(port),
        ],
        cwd=server_working_dir,
        env=server_env,
        stdout=log_file,
        stderr=subprocess.STDOUT,
    )
    base_url = f"http://127.0.0.1:{port}"
    deadline = time.monotonic() + STARTUP_TIMEOUT_SECONDS
    # Use a client with proxy disabled to avoid corporate proxy intercepting localhost
    client = httpx.Client(proxy=None, timeout=2.0)
    try:
        while time.monotonic() < deadline:
            if process.poll() is not None:
                raise RuntimeError(
                    f"server exited during startup:\n{log_path.read_text(errors='replace')}"
                )
            try:
                response = client.get(f"{base_url}/health/ready")
                if response.status_code == 200:
                    break
            except httpx.HTTPError:
                time.sleep(0.3)
        else:
            raise RuntimeError("server did not become ready in time")

        readiness = client.get(f"{base_url}/health/ready").json()
        assert all(
            check["ok"] for check in readiness["checks"].values()
        ), f"server not ready: {readiness}"
        yield base_url
    finally:
        client.close()
        process.terminate()
        try:
            process.wait(timeout=10)
        except subprocess.TimeoutExpired:
            process.kill()
        log_file.close()


@pytest.fixture(scope="session")
def restartable_app_server(
    server_env: dict[str, str],
    migrated_db: str,
    server_working_dir: Path,
    tmp_path_factory: pytest.TempPathFactory,
) -> Iterator[RestartableAppServer]:
    """Start a server whose database/storage survive one deliberate restart."""
    server = RestartableAppServer(
        server_env,
        server_working_dir,
        tmp_path_factory.mktemp("browser-restart-server"),
    )
    server.start()
    try:
        yield server
    finally:
        server.stop()


@pytest.fixture(scope="session")
def session_factory(migrated_db: str) -> Any:
    from migration_intake.persistence.database import (
        create_engine_from_url,
        create_session_factory,
    )

    return create_session_factory(create_engine_from_url(migrated_db))


@pytest.fixture(scope="session")
def catalog_release_id(session_factory: Any) -> str:
    from migration_intake.application.services.catalogs import CatalogPublicationService

    release = CatalogPublicationService(session_factory).get_latest_published()
    assert release is not None, "no published catalog release"
    return str(release["id"])


@pytest.fixture()
def intake(session_factory: Any, catalog_release_id: str, server_env: dict[str, str]) -> dict:
    """
    Create an application (with a CORRELATION identifier) and an open intake.

    Seeded through the service layer so journeys start from a known state
    without paying UI cost. Landmine 3: the identifier is what stops every
    evidence import from quarantining.
    """
    from migration_intake.application.commands import (
        CreateApplicationCommand,
        CreateIntakeCommand,
        IdentifierInput,
    )
    from migration_intake.application.dto import ActorContext
    from migration_intake.application.services.applications import ApplicationService

    actor = ActorContext(
        actor_id=server_env["ACTOR_ID"],
        display_name="Browser Suite Actor",
        actor_type="CONFIGURED",
    )
    service = ApplicationService(session_factory)
    suffix = uuid.uuid4().hex[:6]
    # Correlation identifiers are globally unique (duplicates raise a typed
    # error by design), so each test gets its own rather than sharing one.
    correlation_id = str(uuid.uuid4().int % 90000000 + 10000000)
    application = service.create_application(
        CreateApplicationCommand(
            display_name=f"Browser Suite App {suffix}",
            identifiers=(IdentifierInput(identifier_type="CORRELATION", raw_value=correlation_id),),
            actor=actor,
        )
    )
    created = service.create_intake(
        CreateIntakeCommand(
            application_id=application["id"],
            catalog_release_id=catalog_release_id,
            actor=actor,
        )
    )
    return {
        "application_id": application["id"],
        "intake_id": created["id"],
        "correlation_id": correlation_id,
    }


@pytest.fixture()
def topology_ready_intake(intake: dict, session_factory: Any, server_env: dict[str, str]) -> dict:
    """Provide a confirmed application name/acronym for governed preview capture."""
    from migration_intake.application.dto import ActorContext
    from migration_intake.application.services.answers import AnswerService

    actor = ActorContext(
        actor_id=server_env["ACTOR_ID"],
        display_name="Browser Suite Actor",
        actor_type="CONFIGURED",
    )
    answers = AnswerService(session_factory)
    answers.save_answer(
        intake["intake_id"],
        "CTL-002",
        {"first": "Browser Suite App", "second": "BSA"},
        actor,
        change_reason="Topology browser preview fixture",
    )
    answers.confirm_answer(
        intake["intake_id"],
        "CTL-002",
        actor,
        rationale="Topology browser preview fixture",
    )
    return intake


# ---------------------------------------------------------------------------
# Synthetic evidence builders — never client data (AGENTS.md)
# ---------------------------------------------------------------------------


def build_uaq_workbook(correlation_id: str, *, unique: str | None = None) -> bytes:
    """
    Build a synthetic UAQ workbook using only reviewed, mapped columns.

    ``unique`` defeats evidence content-hash dedup and import idempotency
    (landmine 4) so each test genuinely imports rather than silently
    re-reading an earlier run.
    """
    workbook = Workbook()
    sheet = workbook.active
    # Excel truncates worksheet names at 31 characters; the real export is
    # named this way, so the suite exercises the truncated form on purpose.
    sheet.title = "UAQ - Unified Assessment Questi"
    headers = [
        "Correlation ID",
        "App name",
        "App Acronym",
        "INV9-What is the current operational status of the application?",
        "INV7-Available Environments (Dev/Test/Stage/Prod) that exist to support this workload?",
        "Trace Token",
    ]
    sheet.append(headers)
    sheet.append(
        [
            correlation_id,
            "Browser Suite App",
            "BSA",
            "active",
            "Development, Production",
            unique or uuid.uuid4().hex,
        ]
    )
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def fill_intake_form(content: bytes, answers: dict[str, str]) -> bytes:
    """
    Write answers into an exported intake form's ``Response`` column.

    The form is always obtained from the application's own export endpoint so
    the suite cannot drift from the real header contract.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content))
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    code_column = headers.index("Question Code") + 1
    response_column = headers.index("Response") + 1
    source_column = headers.index("Source File") + 1
    confidence_column = headers.index("Confidence") + 1

    remaining = dict(answers)
    for row in range(2, sheet.max_row + 1):
        code = sheet.cell(row, code_column).value
        if code in remaining:
            sheet.cell(row, response_column).value = remaining.pop(code)
            sheet.cell(row, source_column).value = "browser-suite-synthetic"
            sheet.cell(row, confidence_column).value = 0.9
    assert not remaining, f"questions not present in the exported form: {sorted(remaining)}"

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def timestamped_name(stem: str, extension: str = "xlsx") -> str:
    return f"{stem}-{datetime.now(tz=UTC):%H%M%S}-{uuid.uuid4().hex[:6]}.{extension}"


# ---------------------------------------------------------------------------
# Page fixture with invariants attached
# ---------------------------------------------------------------------------


@pytest.fixture()
def recorder(page: Any, app_server: str) -> Iterator[Any]:
    """Attach console/response recording to the page for the whole test."""
    from tests.browser.invariants import PageRecorder

    instance = PageRecorder()
    instance.attach(page, app_server)
    yield instance
    instance.assert_clean()
