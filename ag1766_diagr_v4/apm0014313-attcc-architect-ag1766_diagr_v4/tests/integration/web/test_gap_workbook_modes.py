"""
Integration tests for gap workbook selection modes and provenance columns.

Verifies:
- GET .../gap-workbook.xlsx?include={unresolved,required,all} row selection
- An unknown include mode is rejected with 400
- Answered rows carry Approved Answer Summary and a blank Response
- The new 15-column header reimports and lands provenance on the candidate
- The legacy 12-column header still reimports (backwards compatibility)
- A malformed Confidence cell is a finding, not a 500
"""

from __future__ import annotations

import uuid
from datetime import datetime, timezone
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from migration_intake.application.services.gap_workbook import GapWorkbookService
from migration_intake.config import Settings
from migration_intake.main import create_app
from migration_intake.persistence.models import (
    Actor,
    AnswerInstance,
    AnswerRevision,
    Application,
    Base,
    CatalogQuestion,
    CatalogRelease,
    CatalogSection,
    Intake,
)
from migration_intake.persistence.models_candidates import Candidate, CandidateFinding
from migration_intake.web.security import generate_csrf_token

# Column positions in the exported sheet (1-based, openpyxl style).
COL_ROW_VERSION = 11
COL_RESPONSE = 12
COL_SOURCE_FILE = 13
COL_SOURCE_LOCATOR = 14
COL_CONFIDENCE = 15


@pytest.fixture
def test_settings(tmp_path) -> Settings:
    """Create test settings with an isolated SQLite database."""
    db_path = tmp_path / "test.db"
    return Settings(
        database_url=f"sqlite:///{db_path}",
        app_env="test",
        evidence_root=str(tmp_path / "evidence"),
        actor_id=str(uuid.uuid4()),
        actor_display_name="Test Actor",
        csrf_secret="test-csrf-secret-key-12345",
    )


@pytest.fixture
def test_app(test_settings):
    """Create test application."""
    return create_app(test_settings)


@pytest.fixture
def client(test_app) -> TestClient:
    """Create test client."""
    return TestClient(test_app)


@pytest.fixture
def session_factory(test_app):
    """Get session factory from app state."""
    return test_app.state.session_factory


@pytest.fixture
def seeded_catalog(test_app, session_factory, test_settings) -> dict:
    """Seed an intake whose pinned catalog mixes required, optional, and inactive.

    Active questions:
      Q-REQ-1 (REQUIRED, answered)
      Q-REQ-2 (REQUIRED, unanswered)
      Q-REQ-3 (REQUIRED, unanswered)
      Q-OPT-1 (OPTIONAL, unanswered)
      Q-OPT-2 (RECOMMENDED, unanswered)
    Plus Q-INACTIVE (REQUIRED but is_active=False), which no mode may export.
    """
    Base.metadata.create_all(test_app.state.engine)
    now = datetime.now(tz=timezone.utc)

    with session_factory() as session:
        actor_id = test_settings.actor_id
        session.add(Actor(id=actor_id, display_name="Test Actor", created_at=now))

        app_id = str(uuid.uuid4())
        session.add(
            Application(
                id=app_id,
                state="ACTIVE",
                display_name="Test App",
                created_at=now,
                updated_at=now,
                row_version=1,
                created_by_id=actor_id,
            )
        )

        catalog_id = str(uuid.uuid4())
        session.add(
            CatalogRelease(
                id=catalog_id,
                semantic_version="1.0.0",
                source_filename="catalog.yaml",
                source_sha256="b" * 64,
                compiler_version="1.0",
                pub_state="PUBLISHED",
                published_at=now,
                created_at=now,
            )
        )

        intake_id = str(uuid.uuid4())
        session.add(
            Intake(
                id=intake_id,
                application_id=app_id,
                catalog_id=catalog_id,
                state="DRAFT",
                created_at=now,
                updated_at=now,
                row_version=1,
                created_by_id=actor_id,
            )
        )

        section_id = str(uuid.uuid4())
        session.add(
            CatalogSection(
                id=section_id,
                release_id=catalog_id,
                section_code="SEC-M",
                display_name="Modes Section",
                display_order=1,
            )
        )
        session.flush()

        question_ids: dict[str, str] = {}
        specs = [
            ("Q-REQ-1", "REQUIRED", True),
            ("Q-REQ-2", "REQUIRED", True),
            ("Q-REQ-3", "REQUIRED", True),
            ("Q-OPT-1", "OPTIONAL", True),
            ("Q-OPT-2", "RECOMMENDED", True),
            ("Q-INACTIVE", "REQUIRED", False),
        ]
        for order, (code, level, active) in enumerate(specs, start=1):
            question_id = str(uuid.uuid4())
            question_ids[code] = question_id
            session.add(
                CatalogQuestion(
                    id=question_id,
                    section_id=section_id,
                    question_code=code,
                    question_text=f"Question {code}?",
                    response_type="TEXT",
                    required_level=level,
                    collection_mode="MANUAL",
                    display_order=order,
                    is_active=active,
                )
            )
        session.flush()

        # Answer Q-REQ-1 so it is required-but-resolved.
        instance_id = str(uuid.uuid4())
        session.add(
            AnswerInstance(
                id=instance_id,
                intake_id=intake_id,
                question_id=question_ids["Q-REQ-1"],
                created_at=now,
                updated_at=now,
                current_rev_id=None,
                applicability="APPLICABLE",
                value_state="SAVED",
                review_state="UNREVIEWED",
                row_version=3,
            )
        )
        session.flush()
        revision_id = str(uuid.uuid4())
        session.add(
            AnswerRevision(
                id=revision_id,
                instance_id=instance_id,
                revision_number=1,
                response_json={"value": "Wave 2"},
                confirm_state="DRAFT",
                authored_at=now,
                authored_by_id=actor_id,
                response_schema_version="1.0",
            )
        )
        session.flush()
        session.execute(
            select(AnswerInstance).where(AnswerInstance.id == instance_id)
        ).scalar_one().current_rev_id = revision_id

        session.commit()

        return {
            "actor_id": actor_id,
            "app_id": app_id,
            "catalog_id": catalog_id,
            "intake_id": intake_id,
            "question_ids": question_ids,
        }


def _export(client: TestClient, seeded: dict, include: str | None = None):
    """GET the gap workbook, optionally with an include mode."""
    url = (
        f"/applications/{seeded['app_id']}"
        f"/intakes/{seeded['intake_id']}/gap-workbook.xlsx"
    )
    params = {} if include is None else {"include": include}
    return client.get(url, params=params)


def _codes(content: bytes) -> list[str]:
    """Read the exported Question Code column."""
    workbook = load_workbook(BytesIO(content), read_only=True)
    rows = list(workbook.active.values)
    return [row[5] for row in rows[1:]]


def _reimport(client: TestClient, seeded: dict, settings: Settings, payload: bytes):
    """POST a completed workbook back for candidate creation."""
    return client.post(
        f"/applications/{seeded['app_id']}"
        f"/intakes/{seeded['intake_id']}/gap-workbook/reimport",
        files={
            "file": (
                "completed-gap.xlsx",
                payload,
                "application/vnd.openxmlformats-officedocument"
                ".spreadsheetml.sheet",
            )
        },
        data={"_csrf_token": generate_csrf_token(settings.csrf_secret)},
    )


class TestIncludeModes:
    """Tests for the include selection modes."""

    def test_default_is_unresolved_required_only(
        self, client: TestClient, seeded_catalog: dict
    ) -> None:
        """No include parameter keeps the original unresolved-required behaviour."""
        response = _export(client, seeded_catalog)

        assert response.status_code == 200
        assert response.headers["x-include-mode"] == "unresolved"
        # Q-REQ-1 is answered; Q-OPT-* are optional; Q-INACTIVE is inactive.
        assert _codes(response.content) == ["Q-REQ-2", "Q-REQ-3"]
        assert response.headers["x-unresolved-question-count"] == "2"

    def test_explicit_unresolved_matches_default(
        self, client: TestClient, seeded_catalog: dict
    ) -> None:
        """Passing include=unresolved is identical to omitting it."""
        default = _export(client, seeded_catalog)
        explicit = _export(client, seeded_catalog, "unresolved")

        assert explicit.status_code == 200
        assert _codes(explicit.content) == _codes(default.content)

    def test_required_mode_includes_answered_required(
        self, client: TestClient, seeded_catalog: dict
    ) -> None:
        """include=required adds the answered required question back in."""
        response = _export(client, seeded_catalog, "required")

        assert response.status_code == 200
        assert response.headers["x-include-mode"] == "required"
        assert _codes(response.content) == ["Q-REQ-1", "Q-REQ-2", "Q-REQ-3"]
        assert response.headers["x-unresolved-question-count"] == "3"

    def test_all_mode_includes_optional_questions(
        self, client: TestClient, seeded_catalog: dict
    ) -> None:
        """include=all exports every active question regardless of level."""
        response = _export(client, seeded_catalog, "all")

        assert response.status_code == 200
        assert response.headers["x-include-mode"] == "all"
        assert _codes(response.content) == [
            "Q-REQ-1",
            "Q-REQ-2",
            "Q-REQ-3",
            "Q-OPT-1",
            "Q-OPT-2",
        ]
        assert response.headers["x-unresolved-question-count"] == "5"

    def test_no_mode_exports_inactive_questions(
        self, client: TestClient, seeded_catalog: dict
    ) -> None:
        """An inactive question is never part of the intake form."""
        for include in ("unresolved", "required", "all"):
            codes = _codes(_export(client, seeded_catalog, include).content)
            assert "Q-INACTIVE" not in codes

    @pytest.mark.parametrize("include", ["", "REQUIRED", "everything", "unresolved "])
    def test_unknown_include_is_rejected(
        self, client: TestClient, seeded_catalog: dict, include: str
    ) -> None:
        """An unknown include value is a 400, never a silent default."""
        response = _export(client, seeded_catalog, include)

        assert response.status_code == 400
        assert "include mode" in response.json()["detail"]

    def test_unknown_include_rejected_before_intake_lookup(
        self, client: TestClient, seeded_catalog: dict
    ) -> None:
        """Mode validation is independent of application ownership checks."""
        response = client.get(
            f"/applications/{uuid.uuid4()}"
            f"/intakes/{seeded_catalog['intake_id']}/gap-workbook.xlsx",
            params={"include": "nonsense"},
        )

        assert response.status_code == 400


class TestApprovedAnswerSummary:
    """Tests for rendering the current answer on already-answered rows."""

    def test_answered_row_shows_summary_and_blank_response(
        self, client: TestClient, seeded_catalog: dict
    ) -> None:
        """The filler sees the current value to confirm, and an empty Response."""
        response = _export(client, seeded_catalog, "required")
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        rows = list(workbook.active.values)
        by_code = {row[5]: row for row in rows[1:]}

        answered = by_code["Q-REQ-1"]
        assert answered[9] == "value=Wave 2"
        assert answered[COL_RESPONSE - 1] in (None, "")
        # The staleness token reflects the existing answer instance.
        assert str(answered[COL_ROW_VERSION - 1]) == "3"

        unanswered = by_code["Q-REQ-2"]
        assert unanswered[9] in (None, "")

    def test_summary_populated_in_all_mode(
        self, client: TestClient, seeded_catalog: dict
    ) -> None:
        """include=all renders the summary too."""
        response = _export(client, seeded_catalog, "all")
        rows = list(load_workbook(BytesIO(response.content), read_only=True).active.values)
        by_code = {row[5]: row for row in rows[1:]}

        assert by_code["Q-REQ-1"][9] == "value=Wave 2"
        assert by_code["Q-REQ-1"][COL_RESPONSE - 1] in (None, "")


class TestProvenanceHeaders:
    """Tests for the appended provenance columns."""

    def test_export_appends_provenance_columns(
        self, client: TestClient, seeded_catalog: dict
    ) -> None:
        """The header ends with the three provenance columns."""
        response = _export(client, seeded_catalog, "all")
        rows = list(load_workbook(BytesIO(response.content), read_only=True).active.values)

        assert tuple(rows[0]) == GapWorkbookService._HEADERS
        assert tuple(rows[0][-3:]) == (
            "Source File",
            "Source Locator",
            "Confidence",
        )
        # Legacy column order is unchanged ahead of them.
        assert tuple(rows[0][:12]) == GapWorkbookService._LEGACY_HEADERS

    def test_reimport_persists_provenance_on_candidate(
        self,
        client: TestClient,
        seeded_catalog: dict,
        session_factory,
        test_settings: Settings,
    ) -> None:
        """Provenance supplied by the filler lands on the created candidate."""
        exported = _export(client, seeded_catalog)
        workbook = load_workbook(BytesIO(exported.content))
        sheet = workbook.active
        sheet.cell(row=2, column=COL_RESPONSE).value = "Wave 1"
        sheet.cell(row=2, column=COL_SOURCE_FILE).value = "migration-tracker.xlsx"
        sheet.cell(row=2, column=COL_SOURCE_LOCATOR).value = "Apps!B14"
        sheet.cell(row=2, column=COL_CONFIDENCE).value = 0.9
        payload = BytesIO()
        workbook.save(payload)

        response = _reimport(
            client, seeded_catalog, test_settings, payload.getvalue()
        )

        assert response.status_code == 200
        assert response.json()["candidate_count"] == 1
        assert response.json()["finding_count"] == 0
        with session_factory() as session:
            candidate = session.execute(
                select(Candidate).where(Candidate.target_key == "Q-REQ-2")
            ).scalar_one()
            assert candidate.state == "PROPOSED"
            assert candidate.confidence == pytest.approx(0.9)
            assert candidate.source_locator["source_file"] == "migration-tracker.xlsx"
            assert candidate.source_locator["source_ref"] == "Apps!B14"
            assert candidate.source_locator["row"] == 2
            assert candidate.origin == (
                "gap_workbook_reimport:migration-tracker.xlsx"
            )

    def test_reimport_without_provenance_keeps_prior_behaviour(
        self,
        client: TestClient,
        seeded_catalog: dict,
        session_factory,
        test_settings: Settings,
    ) -> None:
        """Blank provenance cells leave the candidate exactly as before."""
        exported = _export(client, seeded_catalog)
        workbook = load_workbook(BytesIO(exported.content))
        workbook.active.cell(row=2, column=COL_RESPONSE).value = "Wave 1"
        payload = BytesIO()
        workbook.save(payload)

        response = _reimport(
            client, seeded_catalog, test_settings, payload.getvalue()
        )

        assert response.status_code == 200
        with session_factory() as session:
            candidate = session.execute(
                select(Candidate).where(Candidate.target_key == "Q-REQ-2")
            ).scalar_one()
            assert candidate.origin == "gap_workbook_reimport"
            assert candidate.confidence is None
            assert "source_file" not in candidate.source_locator
            assert "source_ref" not in candidate.source_locator

    def test_legacy_twelve_column_workbook_still_reimports(
        self,
        client: TestClient,
        seeded_catalog: dict,
        session_factory,
        test_settings: Settings,
    ) -> None:
        """Regression guard: workbooks exported before provenance existed import."""
        exported = _export(client, seeded_catalog)
        workbook = load_workbook(BytesIO(exported.content))
        sheet = workbook.active
        sheet.delete_cols(COL_SOURCE_FILE, 3)
        assert tuple(next(sheet.values)) == GapWorkbookService._LEGACY_HEADERS
        sheet.cell(row=2, column=COL_RESPONSE).value = "Wave 1"
        payload = BytesIO()
        workbook.save(payload)

        response = _reimport(
            client, seeded_catalog, test_settings, payload.getvalue()
        )

        assert response.status_code == 200
        assert response.json()["candidate_count"] == 1
        with session_factory() as session:
            candidate = session.execute(
                select(Candidate).where(Candidate.target_key == "Q-REQ-2")
            ).scalar_one()
            assert candidate.state == "PROPOSED"
            assert candidate.origin == "gap_workbook_reimport"
            assert candidate.confidence is None

    def test_answered_row_corrects_as_a_proposal(
        self,
        client: TestClient,
        seeded_catalog: dict,
        session_factory,
        test_settings: Settings,
    ) -> None:
        """A corrected answered row is proposed, never written canonically."""
        exported = _export(client, seeded_catalog, "required")
        workbook = load_workbook(BytesIO(exported.content))
        sheet = workbook.active
        # Row 2 is Q-REQ-1, the already-answered question.
        assert sheet.cell(row=2, column=6).value == "Q-REQ-1"
        sheet.cell(row=2, column=COL_RESPONSE).value = "Wave 3"
        sheet.cell(row=2, column=COL_SOURCE_FILE).value = "kickoff-notes.docx"
        payload = BytesIO()
        workbook.save(payload)

        response = _reimport(
            client, seeded_catalog, test_settings, payload.getvalue()
        )

        assert response.status_code == 200
        with session_factory() as session:
            candidate = session.execute(
                select(Candidate).where(Candidate.target_key == "Q-REQ-1")
            ).scalar_one()
            assert candidate.state == "PROPOSED"
            assert candidate.raw_value_json == {"response": "Wave 3"}
            assert candidate.source_locator["source_file"] == "kickoff-notes.docx"
            # The canonical revision is untouched by the import.
            revision = session.execute(select(AnswerRevision)).scalar_one()
            assert revision.response_json == {"value": "Wave 2"}

    def test_optional_row_is_still_rejected_on_reimport(
        self,
        client: TestClient,
        seeded_catalog: dict,
        session_factory,
        test_settings: Settings,
    ) -> None:
        """Documents a known limit: reimport still targets active REQUIRED only.

        ``include=all`` can export optional questions for completion, but the
        reimport contract's active-required target check is deliberately
        unchanged, so an optional row is refused rather than persisted.
        """
        exported = _export(client, seeded_catalog, "all")
        workbook = load_workbook(BytesIO(exported.content))
        sheet = workbook.active
        optional_row = next(
            row_number
            for row_number in range(2, sheet.max_row + 1)
            if sheet.cell(row=row_number, column=6).value == "Q-OPT-1"
        )
        sheet.cell(row=optional_row, column=COL_RESPONSE).value = "Nice to have"
        payload = BytesIO()
        workbook.save(payload)

        response = _reimport(
            client, seeded_catalog, test_settings, payload.getvalue()
        )

        assert response.status_code == 400
        assert "invalid question" in response.json()["detail"]
        with session_factory() as session:
            assert session.execute(select(Candidate)).scalars().all() == []

    def test_unrelated_header_is_still_rejected(
        self,
        client: TestClient,
        seeded_catalog: dict,
        test_settings: Settings,
    ) -> None:
        """Strict rejection is kept for anything that is neither layout."""
        exported = _export(client, seeded_catalog)
        workbook = load_workbook(BytesIO(exported.content))
        workbook.active.cell(row=1, column=1).value = "Not An Application ID"
        payload = BytesIO()
        workbook.save(payload)

        response = _reimport(
            client, seeded_catalog, test_settings, payload.getvalue()
        )

        assert response.status_code == 400
        assert "headers" in response.json()["detail"]


class TestConfidenceParsing:
    """Tests for defensive Confidence handling."""

    @pytest.mark.parametrize(
        "value", ["very high", "0.9 (est)", "n/a", "1.5", "-0.2", "abc"]
    )
    def test_malformed_confidence_does_not_error(
        self,
        client: TestClient,
        seeded_catalog: dict,
        session_factory,
        test_settings: Settings,
        value: str,
    ) -> None:
        """An unusable Confidence cell is ignored and reported, never a 500."""
        exported = _export(client, seeded_catalog)
        workbook = load_workbook(BytesIO(exported.content))
        sheet = workbook.active
        sheet.cell(row=2, column=COL_RESPONSE).value = "Wave 1"
        sheet.cell(row=2, column=COL_CONFIDENCE).value = value
        payload = BytesIO()
        workbook.save(payload)

        response = _reimport(
            client, seeded_catalog, test_settings, payload.getvalue()
        )

        assert response.status_code == 200
        assert response.json()["candidate_count"] == 1
        assert response.json()["finding_count"] == 1
        with session_factory() as session:
            candidate = session.execute(
                select(Candidate).where(Candidate.target_key == "Q-REQ-2")
            ).scalar_one()
            assert candidate.confidence is None
            finding = session.execute(
                select(CandidateFinding).where(
                    CandidateFinding.candidate_id == candidate.id
                )
            ).scalar_one()
            assert finding.finding_type == "INVALID_CONFIDENCE"
            assert finding.severity == "WARNING"

    def test_confidence_as_text_number_is_accepted(
        self,
        client: TestClient,
        seeded_catalog: dict,
        session_factory,
        test_settings: Settings,
    ) -> None:
        """A numeric string cell is parsed rather than discarded."""
        exported = _export(client, seeded_catalog)
        workbook = load_workbook(BytesIO(exported.content))
        sheet = workbook.active
        sheet.cell(row=2, column=COL_RESPONSE).value = "Wave 1"
        sheet.cell(row=2, column=COL_CONFIDENCE).value = " 0.75 "
        payload = BytesIO()
        workbook.save(payload)

        response = _reimport(
            client, seeded_catalog, test_settings, payload.getvalue()
        )

        assert response.status_code == 200
        with session_factory() as session:
            candidate = session.execute(
                select(Candidate).where(Candidate.target_key == "Q-REQ-2")
            ).scalar_one()
            assert candidate.confidence == pytest.approx(0.75)
