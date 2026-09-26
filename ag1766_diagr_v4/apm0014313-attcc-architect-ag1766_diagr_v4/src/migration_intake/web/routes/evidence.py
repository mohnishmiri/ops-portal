"""
Evidence routes (UI05) — Sources, upload, and import summary.

Routes:
- GET  /applications/{app}/intakes/{intake}/sources
- GET  /applications/{app}/intakes/{intake}/imports/{run}
- GET  /applications/{app}/intakes/{intake}/imports/{run}/coverage
- POST /applications/{app}/intakes/{intake}/imports/{run}/candidates/bulk
- GET  /applications/{app}/intakes/{intake}/imports/{run}/diagnostics.csv
"""

from __future__ import annotations

import csv
import io
import logging
from datetime import datetime
from pathlib import Path
from typing import Annotated, Any
from urllib.parse import urlencode

from fastapi import APIRouter, Depends, Form, HTTPException, Request, status
from fastapi.responses import RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy.orm import Session
from starlette.responses import Response

from migration_intake.application.dto import ActorContext
from migration_intake.application.errors import (
    ConcurrencyConflictError,
)
from migration_intake.application.queries import ApplicationQueryService
from migration_intake.application.services.candidates import (
    CandidateNotFoundError,
    CandidateNotProposedError,
    CandidateService,
    CandidateTargetNotInCatalogError,
    DeferReasonRequiredError,
    RejectReasonRequiredError,
)
from migration_intake.application.services.import_coverage import ImportCoverageService
from migration_intake.persistence.repositories.answers import AnswerRepository
from migration_intake.persistence.repositories.applications import ApplicationRepository
from migration_intake.persistence.repositories.candidates import CandidateRepository
from migration_intake.persistence.repositories.catalogs import CatalogRepository
from migration_intake.persistence.repositories.evidence import EvidenceRepository
from migration_intake.persistence.repositories.imports import ImportRepository
from migration_intake.persistence.repositories.intakes import IntakeRepository
from migration_intake.web.security import (
    CONFIGURED_ACTOR_CAPABILITIES,
    generate_csrf_token,
    validate_csrf_token,
)

router = APIRouter(tags=["Evidence"])
logger = logging.getLogger(__name__)

# Templates
_template_dir = Path(__file__).parent.parent / "templates"
templates = Jinja2Templates(directory=str(_template_dir))

#: Findings a reviewer must act on are shown before informational noise.
_SEVERITY_RANK: dict[str, int] = {"ERROR": 0, "WARNING": 1, "INFO": 2}

_NO_IDENTITY_REMEDIATION = (
    "No application identity could be read from this legacy intake workbook."
)

_NO_IDENTITY_FIX_STEPS = (
    "Confirm the workbook contains the application correlation identifier.",
    "Correct the identifier in the workbook and upload it again.",
)


def _human_value(value: object) -> str:
    """
    Render a stored raw/normalized response payload as reviewer-facing text.

    Presentation only: the underlying JSON stays authoritative and is still
    available verbatim in the diagnostics export.
    """
    if value is None or value == {}:
        return "—"
    if not isinstance(value, dict):
        return str(value)

    if "people" in value and isinstance(value["people"], list):
        rendered = [
            f"{person.get('name') or person.get('email') or person.get('attuid') or '?'}"
            + (" (primary)" if person.get("primary") else "")
            for person in value["people"]
            if isinstance(person, dict)
        ]
        return ", ".join(rendered) or "—"
    if "values" in value and isinstance(value["values"], list):
        return ", ".join(str(item) for item in value["values"]) or "—"
    if "scope" in value:
        direction = value.get("direction")
        return f"{value['scope']} ({direction})" if direction else str(value["scope"])
    for single_key in ("value", "text"):
        if single_key in value and len(value) == 1:
            return str(value[single_key])
    if "first" in value or "second" in value:
        parts = [str(value[key]) for key in ("first", "second") if value.get(key) is not None]
        return " / ".join(parts) or "—"

    return ", ".join(f"{key}: {item}" for key, item in sorted(value.items()) if item is not None)


def _human_locator(locator: object) -> str:
    """Render a source locator as a readable evidence trail."""
    if isinstance(locator, dict):
        parts = [str(locator[key]) for key in ("sheet", "row", "column") if locator.get(key)]
        return " / ".join(parts) if parts else "—"
    return str(locator) if locator else "—"


def _group_findings(findings: list[dict]) -> list[dict]:
    """
    Group findings by type so 200 informational rows collapse to one line.

    Returns action-required severities first; each group carries its own
    findings so the template can disclose them on demand.
    """
    grouped: dict[tuple[str, str], list[dict]] = {}
    for finding in findings:
        key = (str(finding.get("severity", "INFO")), str(finding.get("finding_type", "UNKNOWN")))
        grouped.setdefault(key, []).append(finding)

    return [
        {
            "severity": severity,
            "finding_type": finding_type,
            "count": len(items),
            "findings": items,
            "action_required": severity in ("ERROR", "WARNING"),
        }
        for (severity, finding_type), items in sorted(
            grouped.items(), key=lambda item: (_SEVERITY_RANK.get(item[0][0], 3), item[0][1])
        )
    ]


def _short_datetime(value: object) -> str:
    """Render a timestamp as ``10 Sep, 19:04`` — microseconds help nobody scan a list."""
    if isinstance(value, datetime):
        return f"{value.day} {value:%b}, {value:%H:%M}"
    return str(value) if value else "—"


templates.env.filters["human_value"] = _human_value
templates.env.filters["human_locator"] = _human_locator
templates.env.filters["short_datetime"] = _short_datetime


# ─────────────────────────────────────────────────────────────────────────────
# Import outcome vocabulary
# ─────────────────────────────────────────────────────────────────────────────

#: Plain-language remediation for every non-matching identity verdict. A
#: quarantined import used to end in a bare error, which reads as "your
#: application record is wrong" even when the file is simply the wrong one.
_IDENTITY_REMEDIATION: dict[str, str] = {
    "APPLICATION_MISMATCH": (
        "The correlation identifier written in the file belongs to a different "
        "application. Either you uploaded another application's export, or this "
        "application's correlation identifier is wrong."
    ),
    "APPLICATION_ID_MISSING": (
        "No correlation identifier could be read from the file, so it cannot be "
        "tied to this application."
    ),
    "AMBIGUOUS_APPLICATION_MATCH": (
        "The identifier in the file matches more than one application, so the "
        "import cannot be attributed safely."
    ),
    "SOURCE_ID_CONFLICT": (
        "The file carries more than one correlation identifier, so it describes "
        "more than one application."
    ),
}

#: Ordered remediation steps offered behind the "How to fix" disclosure.
_IDENTITY_FIX_STEPS: dict[str, tuple[str, ...]] = {
    "APPLICATION_MISMATCH": (
        "Open the source file and check the Correlation ID column.",
        "If the file is for another application, upload it under that application.",
        "If the identifier here is wrong, correct it on the application record, "
        "then upload the file again.",
    ),
    "APPLICATION_ID_MISSING": (
        "Confirm the export includes its Correlation ID column and that the row "
        "for this application is populated.",
        "Re-export the source document with the identifier present, then upload "
        "it again.",
    ),
    "AMBIGUOUS_APPLICATION_MATCH": (
        "Two applications share this identifier. Resolve the duplicate on the "
        "application records first.",
        "Upload the file again once exactly one application owns the identifier.",
    ),
    "SOURCE_ID_CONFLICT": (
        "Split the export so each file covers a single application.",
        "Upload each single-application file separately.",
    ),
}

def _identity_summary(run: dict | None, correlation_id: str | None) -> dict[str, object]:
    """
    Describe an import's identity verdict in reviewer-facing language.

    Returns the pill label and tone, a one-line explanation, and the ordered
    remediation steps that back the "How to fix" affordance.
    """
    if run is None:
        return {
            "code": None,
            "label": "Not processed",
            "tone": "neutral",
            "quarantined": False,
            "explanation": "",
            "fix_steps": (),
        }

    decision = run.get("identity_decision")
    if decision == "APPLICATION_MATCHED":
        return {
            "code": decision,
            "label": "Identity verified",
            "tone": "success",
            "quarantined": False,
            "explanation": "",
            "fix_steps": (),
        }

    if decision:
        found = run.get("source_identity_raw")
        located = ""
        if found:
            expected = f" (expected {correlation_id})" if correlation_id else ""
            located = (
                f"Correlation ID {found} in the file does not match this "
                f"application{expected}."
            )
        explanation = _IDENTITY_REMEDIATION.get(
            str(decision), "This import could not be attributed to this application."
        )
        return {
            "code": decision,
            "label": "Quarantined",
            "tone": "danger",
            "quarantined": True,
            "explanation": f"{located} {explanation}".strip(),
            "fix_steps": _IDENTITY_FIX_STEPS.get(str(decision), ()),
        }

    return {
        "code": None,
        "label": "Identity not evaluated",
        "tone": "warning",
        "quarantined": True,
        "explanation": _NO_IDENTITY_REMEDIATION,
        "fix_steps": _NO_IDENTITY_FIX_STEPS,
    }


def _current_answers_by_code(
    session: Session, intake_id: str, catalog_id: str, question_codes: set[str]
) -> dict[str, object]:
    """
    Return the canonical answer currently held for each question code.

    Only codes with a non-empty current revision appear, so a caller can treat
    membership as "accepting this proposal would replace an existing answer".
    """
    catalog = CatalogRepository(session)
    answers = AnswerRepository(session)
    current: dict[str, object] = {}
    for code in question_codes:
        question = catalog.get_question_by_code(catalog_id, code)
        if question is None:
            continue
        instance = answers.get_instance(intake_id, question["id"])
        if instance is None:
            continue
        revision = answers.get_current_revision(instance["id"])
        if revision is None:
            continue
        response = revision["response_json"]
        if response in (None, {}, []):
            continue
        current[code] = response
    return current


# ─────────────────────────────────────────────────────────────────────────────
# File parsing helpers
# ─────────────────────────────────────────────────────────────────────────────


def _parse_csv_to_sheets(content: bytes) -> dict[str, list[dict]]:
    """
    Parse CSV content into sheet_rows format.

    For CSV files, we create a single "UAQ" sheet containing all rows.
    The first non-schema row is treated as headers.

    Returns:
        dict mapping sheet name to list of row dicts
    """
    try:
        # Decode content
        text = content.decode("utf-8-sig")  # Handle BOM if present
    except UnicodeDecodeError:
        text = content.decode("latin-1")

    if not text.strip():
        return {}

    # Skip SharePoint schema row if present (starts with "ListSchema=" or quoted version).
    csv_text = text
    first_line = text.split("\n", 1)[0] if "\n" in text else text

    # Check for unquoted ListSchema= or quoted "ListSchema=...
    if first_line.startswith("ListSchema=") or first_line.startswith('"ListSchema='):
        _, _, csv_text = text.partition("\n")

    reader = csv.DictReader(io.StringIO(csv_text, newline=""))
    if not reader.fieldnames or not any(header.strip() for header in reader.fieldnames):
        return {}
    rows = list(reader)

    if not rows:
        return {}

    # Return as "UAQ" sheet (Unified Assessment Questionnaire)
    return {"UAQ": rows}


def _parse_xlsx_to_sheets(content: bytes) -> dict[str, list[dict]]:
    """
    Parse Excel workbook content into sheet_rows format.

    Returns:
        dict mapping sheet name to list of row dicts
    """
    try:
        import openpyxl
    except ImportError:
        # openpyxl not available
        return {}

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        return {}

    sheet_rows: dict[str, list[dict]] = {}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows_iter = ws.iter_rows(values_only=True)

        # First row is headers
        try:
            headers = next(rows_iter)
        except StopIteration:
            continue

        # SharePoint exports may place a ListSchema metadata row before the
        # actual answer headers.
        if headers and str(headers[0] or "").startswith("ListSchema="):
            try:
                headers = next(rows_iter)
            except StopIteration:
                continue

        if not headers or all(h is None for h in headers):
            continue

        # Convert None headers to empty strings
        headers = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(headers)]

        # Parse data rows
        rows = []
        for row_values in rows_iter:
            if all(v is None for v in row_values):
                continue  # Skip empty rows
            row_dict = {headers[i]: row_values[i] for i in range(min(len(headers), len(row_values)))}
            rows.append(row_dict)

        if rows:
            sheet_rows[sheet_name] = rows

    return sheet_rows


# ─────────────────────────────────────────────────────────────────────────────
# Dependencies
# ─────────────────────────────────────────────────────────────────────────────


def get_actor_context(request: Request) -> ActorContext:
    """Get actor context from app state."""
    settings = request.app.state.settings
    return ActorContext(
        actor_id=settings.actor_id,
        display_name=settings.actor_display_name,
        role_codes=CONFIGURED_ACTOR_CAPABILITIES,
    )


# ─────────────────────────────────────────────────────────────────────────────
# GET /applications/{app}/intakes/{intake}/sources
# ─────────────────────────────────────────────────────────────────────────────


@router.get("/applications/{app_id}/intakes/{intake_id}/sources")
async def list_sources(
    request: Request,
    app_id: str,
    intake_id: str,
    upload_message: str | None = None,
    upload_message_type: str | None = None,
    uploaded_filename: str | None = None,
) -> Response:
    """List legacy intake workbook imports for an intake."""
    session_factory = request.app.state.session_factory
    settings = request.app.state.settings
    csrf_token = generate_csrf_token(request.app.state.settings.csrf_secret)
    questionnaire_status = ApplicationQueryService(session_factory).get_questionnaire_page(
        intake_id
    )

    with session_factory() as session:
        evidence_repo = EvidenceRepository(session)
        import_repo = ImportRepository(session)
        candidate_repo = CandidateRepository(session)

        correlation_id = next(
            (
                identifier["raw_value"]
                for identifier in ApplicationRepository(session).list_identifiers(app_id)
                if identifier["identifier_type"] == "CORRELATION"
            ),
            None,
        )

        # Get evidence items for this intake
        evidence_items = evidence_repo.list_for_intake(intake_id)

        # Enrich with import run info
        sources: list[dict[str, Any]] = []
        for item in evidence_items:
            # Get latest import run for this evidence
            runs = import_repo.list_runs_for_evidence(item["id"])
            latest_run = runs[0] if runs else None

            candidate_count = 0
            proposal_count = 0
            has_reviewable_proposals = False
            problem_count = 0
            if latest_run:
                run_candidates = candidate_repo.get_by_run(latest_run["id"])
                candidate_count = len(run_candidates)
                proposal_count = sum(
                    1
                    for candidate in run_candidates
                    if candidate.state == "PROPOSED" and candidate.target_kind == "QUESTION"
                )
                has_reviewable_proposals = any(
                    candidate.state in {"PROPOSED", "ACCEPTED"}
                    and candidate.target_kind == "QUESTION"
                    for candidate in run_candidates
                )
                problem_count = sum(
                    1
                    for finding in import_repo.list_findings_for_run(latest_run["id"])
                    if finding.get("severity") in ("ERROR", "WARNING")
                )

            sources.append({
                "evidence": item,
                "latest_run": latest_run,
                "candidate_count": candidate_count,
                "proposal_count": proposal_count,
                "has_reviewable_proposals": has_reviewable_proposals,
                "problem_count": problem_count,
                "has_been_processed": latest_run is not None,
                "lane": "Legacy intake",
                "identity": _identity_summary(latest_run, correlation_id),
            })

        # Newest first: a user who has just uploaded looks at the top of the table.
        # str() rather than the raw value: PortableUTC hands back datetimes that
        # may differ in tz-awareness across drivers, and comparing those raises.
        sources.sort(
            key=lambda source: str(source["evidence"]["created_at"] or ""), reverse=True
        )

    return templates.TemplateResponse(
        request=request,
        name="evidence/list.html",
        context={
            "csrf_token": csrf_token,
            "app_id": app_id,
            "intake_id": intake_id,
            "sources": sources,
            "correlation_id": correlation_id,
            "upload_message": upload_message,
            "upload_message_type": upload_message_type or "error",
            "uploaded_filename": uploaded_filename,
            "ai_mapping_available": (
                settings.llm_provider == "mock"
                or (settings.llm_enabled and settings.llm_outbound_enabled)
            ),
            "ai_mapping_provider": settings.llm_provider,
            "ai_mapping_profile": settings.llm_profile,
            "overall_status": (
                questionnaire_status["overall_status"]
                if questionnaire_status
                else {"completed_count": 0, "question_count": 0}
            ),
        },
    )


@router.get("/applications/{app_id}/intakes/{intake_id}/sources/{evidence_id}/delete")
async def delete_source_get(
    app_id: str,
    intake_id: str,
    evidence_id: str,
) -> RedirectResponse:
    """Return safely to Evidence when an old browser tab sends a GET."""
    return RedirectResponse(
        url=f"/applications/{app_id}/intakes/{intake_id}/sources"
        f"?upload_message=Confirm+deletion+from+the+Evidence+page&upload_message_type=error",
        status_code=status.HTTP_303_SEE_OTHER,
    )


@router.post("/applications/{app_id}/intakes/{intake_id}/sources/{evidence_id}/delete")
async def delete_source(
    request: Request,
    app_id: str,
    intake_id: str,
    evidence_id: str,
    csrf_token: Annotated[str, Form(alias="_csrf_token")],
) -> RedirectResponse:
    """Purge an unaccepted Evidence import after explicit confirmation."""
    if not validate_csrf_token(csrf_token, request.app.state.settings.csrf_secret):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Invalid CSRF token")

    try:
        with request.app.state.session_factory() as session:
            deleted = EvidenceRepository(session).purge_import(evidence_id, app_id, intake_id)
            if not deleted:
                raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Evidence not found")
            session.commit()
    except ValueError as exc:
        query = urlencode({"upload_message": str(exc), "upload_message_type": "error"})
        return RedirectResponse(
            url=f"/applications/{app_id}/intakes/{intake_id}/sources?{query}",
            status_code=status.HTTP_303_SEE_OTHER,
        )

    return RedirectResponse(
        url=f"/applications/{app_id}/intakes/{intake_id}/sources?upload_message=Evidence+purged&upload_message_type=success",
        status_code=status.HTTP_303_SEE_OTHER,
    )


# ─────────────────────────────────────────────────────────────────────────────
# GET /applications/{app}/intakes/{intake}/imports/{run}
# ─────────────────────────────────────────────────────────────────────────────


@router.get("/applications/{app_id}/intakes/{intake_id}/imports/{run_id}")
async def import_detail(
    request: Request,
    app_id: str,
    intake_id: str,
    run_id: str,
) -> Response:
    """
    Show import run details.

    Displays sheet outcomes, candidate counts, and findings.
    """
    session_factory = request.app.state.session_factory

    with session_factory() as session:
        import_repo = ImportRepository(session)
        candidate_repo = CandidateRepository(session)

        # Get import run
        run = import_repo.get_run(run_id)
        if run is None or run["application_id"] != app_id or run["intake_id"] != intake_id:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Import run not found",
            )

        # Get sheet results
        sheet_results = import_repo.list_sheet_results_for_run(run_id)

        # Get findings
        findings = import_repo.list_findings_for_run(run_id)

        # Get candidate count
        candidate_count = candidate_repo.count_by_run(run_id)

    return templates.TemplateResponse(
        request=request,
        name="evidence/import_detail.html",
        context={
            "csrf_token": generate_csrf_token(request.app.state.settings.csrf_secret),
            "app_id": app_id,
            "intake_id": intake_id,
            "run": run,
            "sheet_results": sheet_results,
            "findings": findings,
            "finding_groups": _group_findings(findings),
            "candidate_count": candidate_count,
        },
    )


@router.get("/applications/{app_id}/intakes/{intake_id}/imports/{run_id}/coverage")
async def import_coverage(
    request: Request,
    app_id: str,
    intake_id: str,
    run_id: str,
    bulk_message: str | None = None,
    bulk_message_type: str | None = None,
    bulk_detail: str | None = None,
) -> Response:
    """
    Show candidate and remaining-question coverage for an import run.

    Every proposal is rendered beside the answer it would land on, because a
    reviewer previously accepted two proposals targeting the same question and
    the second silently replaced the first. Replacements are flagged here and
    excluded from the default bulk selection.
    """
    session_factory = request.app.state.session_factory
    with session_factory() as session:
        run = ImportRepository(session).get_run(run_id)
    if run is None or run["application_id"] != app_id or run["intake_id"] != intake_id:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Import run not found")

    coverage = ImportCoverageService(session_factory).get_run_coverage(run_id)
    if coverage is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Import coverage not found")

    review_allowed = coverage["identity_decision"] == "APPLICATION_MATCHED"

    with session_factory() as session:
        correlation_id = next(
            (
                identifier["raw_value"]
                for identifier in ApplicationRepository(session).list_identifiers(app_id)
                if identifier["identifier_type"] == "CORRELATION"
            ),
            None,
        )
        intake = IntakeRepository(session).get(intake_id)
        current_answers: dict[str, object] = {}
        if intake is not None:
            current_answers = _current_answers_by_code(
                session,
                intake_id,
                intake["catalog_id"],
                {
                    question["code"] for question in coverage["review_questions"]
                },
            )

    candidates_by_code = {
        candidate["target_key"]: candidate
        for candidate in coverage["candidates"]
        if candidate["target_kind"] == "QUESTION"
        and candidate["state"] in {"PROPOSED", "ACCEPTED"}
    }
    for question in coverage["review_questions"]:
        candidate = candidates_by_code.get(question["code"])
        question["candidate"] = candidate
        question["current_value_json"] = current_answers.get(question["code"])

    reviewable = []
    for candidate in coverage["candidates"]:
        current_value = current_answers.get(candidate["target_key"])
        candidate["current_value_json"] = current_value
        candidate["replaces_existing"] = current_value is not None
        candidate["bulk_eligible"] = bool(
            review_allowed
            and candidate["state"] == "PROPOSED"
            and candidate["target_kind"] == "QUESTION"
            and candidate["is_application_scoped"]
        )
        if candidate["bulk_eligible"]:
            reviewable.append(candidate)

    for question in coverage["review_questions"]:
        candidate = question["candidate"]
        question["candidate_eligible"] = bool(
            candidate and candidate.get("bulk_eligible")
        )

    return templates.TemplateResponse(
        request=request,
        name="evidence/import_coverage.html",
        context={
            "csrf_token": generate_csrf_token(request.app.state.settings.csrf_secret),
            "app_id": app_id,
            "intake_id": intake_id,
            "run": run,
            "coverage": coverage,
            "review_allowed": review_allowed,
            "identity": _identity_summary(run, correlation_id),
            "reviewable_candidates": reviewable,
            "reviewable_count": len(reviewable),
            # Replacements are never pre-selected, so the button's count is the
            # number of additive proposals, not the number of rows.
            "preselected_count": sum(
                1 for candidate in reviewable if not candidate["replaces_existing"]
            ),
            "replacement_count": sum(
                1 for candidate in reviewable if candidate["replaces_existing"]
            ),
            "bulk_message": bulk_message,
            "bulk_message_type": bulk_message_type or "info",
            "bulk_detail": [line for line in (bulk_detail or "").split("|") if line],
        },
    )


# ─────────────────────────────────────────────────────────────────────────────
# POST /applications/{app}/intakes/{intake}/imports/{run}/candidates/bulk
# ─────────────────────────────────────────────────────────────────────────────

#: Per-row failure vocabulary. Every one of these is a genuine outcome for a
#: single candidate, so a batch reports them instead of failing as a whole.
_SKIP_REASONS: dict[type[Exception], str] = {
    ConcurrencyConflictError: "stale — the proposal changed after this page was loaded",
    CandidateNotProposedError: "stale — already decided after this page was loaded",
    CandidateTargetNotInCatalogError: "its question is not in this intake's catalog",
    CandidateNotFoundError: "no longer available",
}

#: Detail lines shown inline; the rest are summarised as a remainder.
_MAX_DETAIL_LINES = 8


def _parse_selection(raw: str) -> tuple[str, int] | None:
    """
    Split a ``candidate_id:row_version`` checkbox value.

    The row version travels with the checkbox because optimistic concurrency
    compares revision numbers; a hardcoded or omitted token silently accepts a
    candidate the reviewer never saw.
    """
    candidate_id, separator, version = raw.partition(":")
    if not separator or not candidate_id:
        return None
    try:
        return candidate_id, int(version)
    except ValueError:
        return None


@router.post("/applications/{app_id}/intakes/{intake_id}/imports/{run_id}/candidates/bulk")
async def bulk_decide_import_candidates(
    request: Request,
    app_id: str,
    intake_id: str,
    run_id: str,
    actor: Annotated[ActorContext, Depends(get_actor_context)],
    csrf_token: Annotated[str, Form(alias="_csrf_token")],
    decision: Annotated[str, Form()],
    selected: Annotated[list[str], Form()] = [],  # noqa: B006 - FastAPI form default
    reason: Annotated[str, Form()] = "",
) -> RedirectResponse:
    """
    Accept or reject many proposals in one action, reporting per-row outcomes.

    Clearing an import used to cost one page round-trip per proposal. This
    route drives the same ``CandidateService`` per candidate, so nothing here
    bypasses the approval rules, and reports exactly how many succeeded rather
    than failing the batch or claiming a success it did not achieve.
    """
    if not validate_csrf_token(csrf_token, request.app.state.settings.csrf_secret):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Invalid CSRF token")

    if decision not in ("accept", "reject"):
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Unknown bulk decision",
        )

    coverage_url = (
        f"/applications/{app_id}/intakes/{intake_id}/imports/{run_id}/coverage"
    )

    def redirect_with(message: str, message_type: str, detail: list[str]) -> RedirectResponse:
        query = urlencode(
            {
                "bulk_message": message,
                "bulk_message_type": message_type,
                "bulk_detail": "|".join(detail),
            }
        )
        return RedirectResponse(
            url=f"{coverage_url}?{query}", status_code=status.HTTP_303_SEE_OTHER
        )

    with request.app.state.session_factory() as session:
        run = ImportRepository(session).get_run(run_id)
    if run is None or run["application_id"] != app_id or run["intake_id"] != intake_id:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Import run not found")
    if run["identity_decision"] != "APPLICATION_MATCHED":
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Import identity is not approved for review",
        )

    if not selected:
        return redirect_with(
            "Nothing was selected, so nothing changed. Tick the proposals you want "
            "to decide first.",
            "error",
            [],
        )
    if decision == "reject" and not reason.strip():
        return redirect_with(
            "Rejecting requires a reason. Nothing was changed.", "error", []
        )

    service = CandidateService(
        request.app.state.session_factory, request.app.state.answer_service
    )

    decided = 0
    skipped: list[str] = []
    for raw in selected:
        parsed = _parse_selection(raw)
        if parsed is None:
            skipped.append("A selected row was malformed and was left untouched")
            continue
        candidate_id, row_version = parsed

        candidate = service.get_candidate(candidate_id)
        label = candidate["target_key"] if candidate else candidate_id[:8]
        if (
            candidate is None
            or candidate["import_run_id"] != run_id
            or candidate["application_id"] != app_id
            or candidate["intake_id"] != intake_id
            or candidate["target_kind"] != "QUESTION"
            or (candidate.get("scope_json") or {}).get("scope") != "APPLICATION"
        ):
            skipped.append(f"{label} — not reviewable from this import")
            continue

        try:
            if decision == "accept":
                service.accept_candidate(candidate_id, actor, expected_version=row_version)
            else:
                service.reject_candidate(
                    candidate_id, reason, actor, expected_version=row_version
                )
        except tuple(_SKIP_REASONS) as error:
            skipped.append(
                f"{label} — {_SKIP_REASONS.get(type(error), 'could not be saved')}"
            )
            continue
        except Exception:
            logger.exception(
                "Bulk %s failed for candidate %s in run %s", decision, candidate_id, run_id
            )
            skipped.append(f"{label} — could not be saved")
            continue
        decided += 1

    verb = "accepted" if decision == "accept" else "rejected"
    if not skipped:
        return redirect_with(f"{decided} {verb}.", "success", [])

    detail = skipped[:_MAX_DETAIL_LINES]
    if len(skipped) > _MAX_DETAIL_LINES:
        detail.append(f"…and {len(skipped) - _MAX_DETAIL_LINES} more")
    return redirect_with(
        f"{decided} {verb}, {len(skipped)} skipped.",
        "success" if decided else "error",
        detail,
    )


@router.post(
    "/applications/{app_id}/intakes/{intake_id}/imports/{run_id}/candidates/{candidate_id}/accept"
)
async def accept_import_candidate(
    request: Request,
    app_id: str,
    intake_id: str,
    run_id: str,
    candidate_id: str,
    actor: Annotated[ActorContext, Depends(get_actor_context)],
    csrf_token: Annotated[str, Form(alias="_csrf_token")],
    row_version: Annotated[int, Form(alias="_row_version")],
) -> RedirectResponse:
    """Accept one imported candidate into the canonical questionnaire answer."""
    if not validate_csrf_token(csrf_token, request.app.state.settings.csrf_secret):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Invalid CSRF token")

    candidate_service = CandidateService(
        request.app.state.session_factory,
        request.app.state.answer_service,
    )
    try:
        candidate = candidate_service.get_candidate(candidate_id)
        if (
            candidate is None
            or candidate["import_run_id"] != run_id
            or candidate["application_id"] != app_id
            or candidate["intake_id"] != intake_id
            or candidate["target_kind"] != "QUESTION"
            or candidate.get("scope_json", {}).get("scope") != "APPLICATION"
        ):
            raise CandidateNotFoundError(f"Candidate {candidate_id} not found")
        with request.app.state.session_factory() as session:
            run = ImportRepository(session).get_run(run_id)
        if run is None or run["identity_decision"] != "APPLICATION_MATCHED":
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Import identity is not approved for review")
        candidate_service.accept_candidate(
            candidate_id,
            actor,
            expected_version=row_version,
        )
    except CandidateNotFoundError as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(exc)) from exc
    except CandidateNotProposedError as exc:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exc)) from exc
    except CandidateTargetNotInCatalogError as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail=str(exc)
        ) from exc
    except ConcurrencyConflictError as exc:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exc)) from exc

    return RedirectResponse(
        url=f"/applications/{app_id}/intakes/{intake_id}/imports/{run_id}/coverage",
        status_code=status.HTTP_303_SEE_OTHER,
    )


@router.post("/applications/{app_id}/intakes/{intake_id}/imports/{run_id}/candidates/{candidate_id}/reject")
async def reject_import_candidate(
    request: Request,
    app_id: str,
    intake_id: str,
    run_id: str,
    candidate_id: str,
    actor: Annotated[ActorContext, Depends(get_actor_context)],
    csrf_token: Annotated[str, Form(alias="_csrf_token")],
    row_version: Annotated[int, Form(alias="_row_version")],
    reason: Annotated[str, Form()],
) -> RedirectResponse:
    """Reject a matched import candidate without changing canonical answers."""
    return await _decide_import_candidate(
        request, app_id, intake_id, run_id, candidate_id, actor, csrf_token, row_version, reason, "reject"
    )


@router.post("/applications/{app_id}/intakes/{intake_id}/imports/{run_id}/candidates/{candidate_id}/defer")
async def defer_import_candidate(
    request: Request,
    app_id: str,
    intake_id: str,
    run_id: str,
    candidate_id: str,
    actor: Annotated[ActorContext, Depends(get_actor_context)],
    csrf_token: Annotated[str, Form(alias="_csrf_token")],
    row_version: Annotated[int, Form(alias="_row_version")],
    reason: Annotated[str, Form()],
) -> RedirectResponse:
    """Defer a matched import candidate without changing canonical answers."""
    return await _decide_import_candidate(
        request, app_id, intake_id, run_id, candidate_id, actor, csrf_token, row_version, reason, "defer"
    )


async def _decide_import_candidate(
    request: Request,
    app_id: str,
    intake_id: str,
    run_id: str,
    candidate_id: str,
    actor: ActorContext,
    csrf_token: str,
    row_version: int,
    reason: str,
    decision: str,
) -> RedirectResponse:
    if not validate_csrf_token(csrf_token, request.app.state.settings.csrf_secret):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Invalid CSRF token")
    service = CandidateService(request.app.state.session_factory, request.app.state.answer_service)
    candidate = service.get_candidate(candidate_id)
    with request.app.state.session_factory() as session:
        run = ImportRepository(session).get_run(run_id)
    if (
        candidate is None
        or candidate["import_run_id"] != run_id
        or candidate["application_id"] != app_id
        or candidate["intake_id"] != intake_id
        or run is None
    ):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Candidate not found")
    if run["identity_decision"] != "APPLICATION_MATCHED":
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Import identity is not approved for review")
    try:
        if decision == "reject":
            service.reject_candidate(candidate_id, reason, actor, expected_version=row_version)
        else:
            service.defer_candidate(candidate_id, reason, actor, expected_version=row_version)
    except (CandidateNotProposedError, ConcurrencyConflictError) as exc:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exc)) from exc
    except (RejectReasonRequiredError, DeferReasonRequiredError) as exc:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail=str(exc)) from exc
    return RedirectResponse(
        url=f"/applications/{app_id}/intakes/{intake_id}/imports/{run_id}/coverage",
        status_code=status.HTTP_303_SEE_OTHER,
    )


# ─────────────────────────────────────────────────────────────────────────────
# GET /applications/{app}/intakes/{intake}/imports/{run}/diagnostics.csv
# ─────────────────────────────────────────────────────────────────────────────


@router.get("/applications/{app_id}/intakes/{intake_id}/imports/{run_id}/diagnostics.csv")
async def diagnostics_csv(
    request: Request,
    app_id: str,
    intake_id: str,
    run_id: str,
) -> StreamingResponse:
    """
    Export import diagnostics as CSV.

    Contains safe codes, sheet/row/cell locators, and bounded messages.
    Applies spreadsheet-formula injection protection.
    """
    session_factory = request.app.state.session_factory

    with session_factory() as session:
        import_repo = ImportRepository(session)

        # Get import run
        run = import_repo.get_run(run_id)
        if run is None:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Import run not found",
            )

        # Get findings
        findings = import_repo.list_findings_for_run(run_id)

    # Generate CSV with formula injection protection
    output = io.StringIO()
    writer = csv.writer(output)

    # Header row
    writer.writerow([
        "Finding Type",
        "Severity",
        "Sheet",
        "Source Locator",
        "Message",
    ])

    # Data rows with formula injection protection
    for finding in findings:
        writer.writerow([
            _sanitize_csv_value(finding.get("finding_type", "")),
            _sanitize_csv_value(finding.get("severity", "")),
            _sanitize_csv_value(finding.get("sheet_name", "")),
            _sanitize_csv_value(finding.get("source_locator", "")),
            _sanitize_csv_value(_truncate_message(finding.get("detail", ""))),
        ])

    output.seek(0)

    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={
            "Content-Disposition": f"attachment; filename=diagnostics_{run_id[:8]}.csv"
        },
    )


def _sanitize_csv_value(value: str) -> str:
    """
    Sanitize a value for CSV export to prevent formula injection.

    Prefixes values starting with =, +, -, @, or tab with a single quote.
    """
    if not value:
        return ""

    value_str = str(value)

    # Formula injection protection
    if value_str and value_str[0] in ("=", "+", "-", "@", "\t"):
        return "'" + value_str

    return value_str


def _truncate_message(message: str, max_length: int = 500) -> str:
    """Truncate message to a bounded length."""
    if not message:
        return ""

    message_str = str(message)
    if len(message_str) > max_length:
        return message_str[:max_length - 3] + "..."

    return message_str
